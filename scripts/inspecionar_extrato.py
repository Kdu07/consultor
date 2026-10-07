"""
Inspeção de extratos XLSX do BTG — somente leitura, sem dados pessoais na saída.

Para que serve: mapear o vocabulário que o BTG usa nos lançamentos (TED, PIX, liquidação de
bolsa, aplicação e resgate de Tesouro, IR, custódia...) antes de escrever as regras do
classificador — docs/PLANO_HISTORICO.md, Bloco 0.

O extrato traz nome, CPF e número da conta. Por isso este script:
  - da Capa, imprime só o período;
  - nunca imprime uma descrição inteira: agrega os lançamentos por um prefixo curto, cortado
    no primeiro separador ou token com dígito, e em transferências mantém só palavras de um
    vocabulário fixo (o nome da contraparte vem logo depois de "PIX RECEBIDO");
  - não imprime o nome do arquivo (o do BTG é o número da conta);
  - não grava nada em disco.

Uso:
    .venv\\Scripts\\python.exe scripts/inspecionar_extrato.py uploads/<arquivo>.xlsx [...]
    .venv\\Scripts\\python.exe scripts/inspecionar_extrato.py            (todos de uploads/)
"""
import sys
from collections import defaultdict
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from app.tools.btg_xlsx_parser import (  # noqa: E402
    ExtratoParseError,
    _aba,
    _cel,
    _data,
    _idx,
    _norm,
    _num,
    _parse_capa,
    _tabelas,
    _txt,
)

# Linha de transferência: o que vem depois da direção costuma ser nome/CPF/conta.
_MARCAS_TRANSFERENCIA = ("pix", "ted", "doc", "transf")
_VOCAB_TRANSFERENCIA = {
    "pix", "ted", "doc", "transf", "transf.", "transferencia", "recebido", "recebida",
    "enviado", "enviada", "devolvido", "devolvida", "devolucao", "estorno", "entre", "contas",
    "conta", "mesma", "titularidade", "custodia", "mercados", "investimento", "investimentos",
    "corrente", "credito", "debito", "para", "da", "do", "de", "em", "cc", "ci", "agendado",
    "agendada", "programado", "programada",
}
_SEPARADORES = {"-", "–", "—", "|", ":", "+"}


def prefixo(descricao: str, max_tokens: int = 3) -> str:
    """Prefixo genérico da descrição — nunca o texto inteiro."""
    tokens: list[str] = []
    for tok in _norm(descricao).split():
        if tok in _SEPARADORES or any(ch.isdigit() for ch in tok):
            break
        tokens.append(tok)
        if len(tokens) == max_tokens:
            break
    if tokens and tokens[0].startswith(_MARCAS_TRANSFERENCIA):
        seguros = []
        for tok in tokens:
            if tok not in _VOCAB_TRANSFERENCIA and not tok.startswith(_MARCAS_TRANSFERENCIA):
                break
            seguros.append(tok)
        tokens = seguros
    return " ".join(tokens) or "(sem texto)"


def _brl(v: float) -> str:
    return f"{v:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")


def inspecionar(caminho: Path, indice: int) -> None:
    conteudo = caminho.read_bytes()
    print(f"\n{'=' * 72}\nExtrato #{indice}  ({len(conteudo) // 1024} KB)")
    wb = load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
    try:
        try:
            data_ref, meta = _parse_capa(wb)
        except ExtratoParseError as e:
            print(f"  Capa ilegível: {e}")
            return
        inicio = meta.get("periodo_inicio")
        print(f"Período: {inicio} → {data_ref}")
        print("Abas:", ", ".join(s for s in wb.sheetnames if _norm(s) != "capa"))

        # ── Sumário: só os nomes de mercado e as datas das colunas ──────────────
        ws = _aba(wb, "Sumario")
        if ws is not None:
            for tab in _tabelas(ws):
                i = _idx(tab, "mercados")
                if i is None:
                    continue
                datas = sorted({d for d in (_data(h) for h in tab.header_raw) if d})
                mercados = [_txt(_cel(r, i)) for r in tab.linhas if _txt(_cel(r, i))]
                print(f"Sumário: datas {datas}; mercados {mercados}")

        # ── Títulos e cabeçalhos de toda aba de dados ───────────────────────────
        for nome in ("Renda Variavel", "Renda Fixa", "Conta Corrente", "Valores em Transito"):
            ws = _aba(wb, nome)
            if ws is None:
                print(f"\n[{nome}] ausente")
                continue
            print(f"\n[{nome}]")
            for tab in _tabelas(ws):
                cab = [h for h in tab.header_raw if h]
                print(f"  • {tab.titulo or '(sem título)'} — {len(tab.linhas)} linha(s); "
                      f"{len(tab.totais)} total(is)")
                print(f"      colunas: {cab}")

        # ── Razão da conta corrente: prefixos agregados ─────────────────────────
        ws = _aba(wb, "Conta Corrente")
        if ws is not None:
            for tab in _tabelas(ws):
                i_desc = _idx(tab, "descricao")
                i_mov = _idx(tab, "movimentacao")
                i_saldo = _idx(tab, "saldo")
                if i_desc is None or i_mov is None:
                    continue
                grupos: dict[str, list[float]] = defaultdict(list)
                saldo_inicial = saldo_final = None
                soma = 0.0
                sem_valor = 0
                for row in tab.linhas:
                    desc = _txt(_cel(row, i_desc))
                    valor = _num(_cel(row, i_mov))
                    saldo = _num(_cel(row, i_saldo))
                    if _norm(desc).startswith("saldo anterior"):
                        saldo_inicial = saldo
                        continue
                    if saldo is not None:
                        saldo_final = saldo
                    if valor is None:
                        sem_valor += 1
                        continue
                    soma += valor
                    grupos[prefixo(desc)].append(valor)
                print(f"\nRazão da conta ('{tab.titulo}'): {sum(len(v) for v in grupos.values())} "
                      f"lançamento(s), {sem_valor} linha(s) sem valor")
                for pref, valores in sorted(grupos.items(), key=lambda kv: -len(kv[1])):
                    cred = sum(v for v in valores if v > 0)
                    deb = sum(v for v in valores if v < 0)
                    print(f"    {len(valores):3d}×  {pref:<34} créditos {_brl(cred):>12}  "
                          f"débitos {_brl(deb):>12}")
                if saldo_inicial is not None and saldo_final is not None:
                    dif = round(saldo_inicial + soma - saldo_final, 2)
                    print(f"    conferência: saldo inicial + Σ − saldo final = {_brl(dif)} "
                          f"({'ok' if abs(dif) <= 0.05 else 'NÃO FECHA'})")
                for tot in tab.totais:
                    rotulo = next((_txt(c) for c in tot if isinstance(c, str) and c.strip()), "")
                    valor = next((_num(c) for c in tot[2:] if _num(c) is not None), None)
                    print(f"    {rotulo}: {_brl(valor) if valor is not None else '-'}")

        # ── Renda variável: transações distintas por bloco de movimentação ──────
        ws = _aba(wb, "Renda Variavel")
        if ws is not None:
            print()
            for tab in _tabelas(ws):
                if not _norm(tab.titulo).startswith("movimentacao"):
                    continue
                i_tr = _idx(tab, "transacao")
                if i_tr is None:
                    continue
                contagem: dict[str, int] = defaultdict(int)
                for row in tab.linhas:
                    contagem[_txt(_cel(row, i_tr)) or "(vazio)"] += 1
                print(f"Transações em '{tab.titulo}': {dict(contagem)}")

        # ── Renda fixa: lotes e lotes adquiridos no período ─────────────────────
        ws = _aba(wb, "Renda Fixa")
        if ws is not None:
            for tab in _tabelas(ws):
                if not _norm(tab.titulo).startswith("detalhamento"):
                    continue
                i_aq = _idx(tab, "aquisicao")
                aquisicoes = [_data(_cel(r, i_aq)) for r in tab.linhas] if i_aq is not None else []
                no_periodo = [a for a in aquisicoes if a and inicio and data_ref and inicio <= a <= data_ref]
                print(f"Lotes em '{tab.titulo}': {len(tab.linhas)} (adquiridos no período: {len(no_periodo)})")

        # ── Valores em trânsito: prefixos ───────────────────────────────────────
        ws = _aba(wb, "Valores em Transito")
        if ws is not None:
            prefs: dict[str, int] = defaultdict(int)
            for tab in _tabelas(ws):
                i_desc = _idx(tab, "descricao")
                if i_desc is None:
                    continue
                for row in tab.linhas:
                    prefs[prefixo(_txt(_cel(row, i_desc)))] += 1
            if prefs:
                print(f"Prefixos em Valores em Trânsito: {dict(prefs)}")
    finally:
        wb.close()


def main() -> int:
    caminhos = [Path(a) for a in sys.argv[1:]] or sorted(Path("uploads").glob("*.xlsx"))
    if not caminhos:
        print("Nenhum XLSX informado e nenhum em uploads/.")
        return 1
    for i, caminho in enumerate(caminhos, 1):
        if not caminho.exists():
            print(f"\nExtrato #{i}: arquivo não encontrado")
            continue
        inspecionar(caminho, i)
    print("\nRevise a saída antes de colar em qualquer lugar: ela não deve ter nome, CPF nem conta.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
