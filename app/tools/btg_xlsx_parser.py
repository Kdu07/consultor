"""
Parser do extrato XLSX da conta de investimento do BTG Pactual.

Fonte única de verdade das posições (o caminho antigo — colagem do texto do PDF —
foi aposentado; ver docs/PLANO_XLSX.md).

Abas lidas: Capa (data de referência), Sumario (checksum), Renda Variavel, Renda Fixa,
Fundos, CriptoAtivos, Conta Corrente (saldo e razão de movimentações), Valores em Trânsito.

Versão 2 do payload (docs/PLANO_HISTORICO.md, Bloco 1): além das posições, o parser entrega
o razão da conta corrente (`lancamentos_conta` — é dele que saem aportes e resgates para a
rentabilidade), os lotes de renda fixa com data de aquisição (`lotes_rf`) e as movimentações
de RV com classe e operação. Descrições de texto livre saem sanitizadas — ver "Privacidade".

Versão 3 do payload (plano "Confiabilidade dos extratos"): abas Fundos e CriptoAtivos viram
posições (classes FUNDO e CRIPTO) e movimentações; códigos decorados com '*' (extrato de
período aberto) saem limpos; nada é descartado em silêncio (`_tabelas` v2); as linhas de
total de cada bloco ficam em `totais_abas` e os fatos atípicos em `avisos_parser` — é sobre
esses dois campos que app/tools/extrato_validacao.py confere as invariantes do extrato.

Princípios do parser (guardrails §1 do PLANO_XLSX):
  - Nada é localizado por índice fixo de linha/coluna: blocos são achados pelo título e
    as colunas são mapeadas pelo cabeçalho normalizado. O layout muda com o conteúdo do
    mês (um mês sem FII não tem o bloco de FIIs).
  - Linha de posição que não parseia NUNCA é silenciada: vai para `linhas_ignoradas`.
  - O total parseado é conferido contra a aba Sumario. Divergência vira aviso, não erro.
  - Sem data de referência não há import: `ExtratoParseError` (nunca inventar a data).
"""
from __future__ import annotations

import logging
import re
import unicodedata
from collections import Counter
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from io import BytesIO
from typing import Any, Optional, Sequence

from openpyxl import load_workbook

logger = logging.getLogger(__name__)

# Tolerância do checksum contra a aba Sumario. O próprio extrato arredonda: em extratos
# reais o Saldo Bruto de um mercado no Sumário difere em centavos da soma das linhas de
# posição. Erro real de parser — uma posição perdida — é de ordem de grandeza muito maior
# (a menor posição de uma carteira é de centenas de reais), então R$ 1,00 separa ruído de
# erro com folga.
TOLERANCIA_CHECKSUM = 1.00

# Conferência do razão da conta corrente (saldo inicial + Σ movimentações = saldo final).
# Diferente do Sumário, o razão é linha a linha e fecha no centavo nos extratos mensais.
# Folga só para arredondamento de float e para microdivergências do próprio BTG.
TOLERANCIA_CONTA = 0.05

# Formato do payload que o parser produz e que fica arquivado em ExtratoImportado. Payload
# sem a chave `versao_parser` é v1: não tem o razão da conta nem os lotes de RF, então o
# motor de desempenho não consegue separar aporte de rendimento naquele mês. O v3 (acima,
# no docstring) acrescenta Fundos/Cripto, `totais_abas` e `avisos_parser`.
VERSAO_PARSER = 3

# Classes de Renda Variável, por título do bloco ("Posição > Ações" etc.)
_CLASSE_RV: dict[str, str] = {
    "acoes": "ACAO",
    "etf": "ETF",
    "fundos listados": "FII",
    "bdr": "BDR",
}

_TESOURO_NOME: dict[str, str] = {
    "LFT": "Tesouro Selic",
    "LTN": "Tesouro Prefixado",
    "NTNB-P": "Tesouro IPCA+",
    "NTNB": "Tesouro IPCA+ com Juros Semestrais",
    "NTN-B": "Tesouro IPCA+ com Juros Semestrais",
    "NTN-F": "Tesouro Prefixado com Juros Semestrais",
}

# Transações da aba de movimentação que são renda (e não compra/venda)
_PROVENTO_MARCAS = ("juros", "rendimento", "dividendo", "provento", "amortizacao")


class ExtratoParseError(Exception):
    """Erro estrutural no arquivo — a tool converte em tool_error para o usuário."""


# ---------------------------------------------------------------------------
# Conversores tolerantes
# ---------------------------------------------------------------------------

def _norm(v: Any) -> str:
    """Minúsculo, sem acento, com whitespace colapsado — para casar rótulos."""
    if v is None:
        return ""
    s = unicodedata.normalize("NFKD", str(v))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().lower()


def _txt(v: Any) -> str:
    """Texto exibível com whitespace colapsado ('BRASIL      ON  NM' → 'BRASIL ON NM')."""
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return _data(v) or ""
    return re.sub(r"\s+", " ", str(v)).strip()


def _num(v: Any) -> Optional[float]:
    """
    float | None. Aceita número nativo do openpyxl, '-' (nulo do BTG), vazio e string
    no formato brasileiro ('1.234,56') — o BTG pode voltar a exportar texto a qualquer
    release.
    """
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, (datetime, date)):
        return None
    s = str(v).strip().replace("R$", "").replace("%", "").strip()
    if not s or s == "-":
        return None
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _data(v: Any) -> Optional[str]:
    """ISO YYYY-MM-DD. Aceita datetime/date do openpyxl e texto dd/mm/aa(aa)."""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if v is None:
        return None
    m = re.search(r"(\d{2})/(\d{2})/(\d{2,4})", str(v))
    if not m:
        return None
    d, mo, y = m.groups()
    if len(y) == 2:
        y = "20" + y
    try:
        return date(int(y), int(mo), int(d)).isoformat()
    except ValueError:
        return None


def _ano(iso: Optional[str]) -> str:
    return iso[:4] if iso else "?"


def _slug(texto: str) -> str:
    """Fragmento de chave a partir de texto livre: sem acento, sem pontuação, maiúsculo."""
    return re.sub(r"[^a-z0-9]+", "-", _norm(texto)).strip("-").upper()


# Ticker da B3: 4 letras + 1-2 dígitos + sufixo opcional (BBAS3, IVVB11, TAEE11, BBAS3F).
_RE_TICKER_B3 = re.compile(r"^[A-Z]{4}\d{1,2}[A-Z]?$")


def _limpar_codigo(v: Any) -> tuple[str, bool]:
    """
    (código limpo, veio decorado?). Em extrato de período aberto o BTG decora os códigos
    com '*' ('BBAS3* ', 'LFT*' — posição movimentada no período). A decoração não pode
    entrar na chave_externa — 'B3:BBAS3*' não casa com 'B3:BBAS3' e o upsert duplicaria a
    posição (aconteceu em produção) — nem no lookup de _TESOURO_NOME ('LFT*' perderia o
    nome do título). Só '*' e espaço saem: siglas de RF privada podem ter hífen e dígitos.
    """
    bruto = _txt(v)
    limpo = re.sub(r"\s+", " ", bruto.replace("*", "")).strip().upper()
    return limpo, bool(bruto) and limpo != bruto.upper()


def _avisar(avisos: Optional[list], tipo: str, aba: str, detalhe: str) -> None:
    """
    Registra um aviso tipado do parser ({tipo, aba, detalhe}) sem duplicar — um extrato de
    período aberto decora TODOS os códigos, e um aviso por linha afogaria os demais.
    """
    if avisos is None:
        return
    item = {"tipo": tipo, "aba": aba, "detalhe": detalhe}
    if item not in avisos:
        avisos.append(item)


def _derivar_quantidade(valor: float, preco: Optional[float]) -> tuple[float, str]:
    """
    Quantidade ilegível: deriva de valor/preço quando há preço; 0.0 é o último recurso
    (o validador promove a erro quando a classe precisa de quantidade para cotação ao
    vivo). Retorna (quantidade, tipo do aviso a registrar).
    """
    if preco:
        return round(valor / preco, 6), "quantidade_derivada"
    return 0.0, "quantidade_ilegivel"


# ---------------------------------------------------------------------------
# Modelos de saída
# ---------------------------------------------------------------------------

@dataclass
class PosicaoParsed:
    ticker: Optional[str]
    nome: str
    classe: str                       # ACAO | ETF | FII | BDR | TESOURO | RF | CAIXA
    quantidade: float
    preco_medio: Optional[float]      # custo médio de aquisição (nunca o preço atual)
    valor_mercado: float              # saldo bruto na data do extrato
    preco_fechamento: Optional[float]
    as_of: Optional[str]              # YYYY-MM-DD
    vencimento: Optional[str] = None  # renda fixa
    taxa_contratada: Optional[str] = None
    custo_total: Optional[float] = None
    # Identidade do papel entre extratos — ver a seção "Identidade do papel entre
    # extratos". O `nome` não serve: o BTG o reescreve de um mês para o outro
    # ('FII HGCR PAXCI' → 'FII HGCR PAXCI ER').
    chave_externa: Optional[str] = None

    def to_dict(self) -> dict:
        return asdict(self)


@dataclass
class ExtratoParsed:
    posicoes: list[PosicaoParsed] = field(default_factory=list)
    data_referencia: Optional[str] = None
    sumario: dict = field(default_factory=dict)
    proventos: list[dict] = field(default_factory=list)
    movimentacoes: list[dict] = field(default_factory=list)
    aluguel: list[dict] = field(default_factory=list)
    valores_em_transito: list[dict] = field(default_factory=list)
    linhas_ignoradas: list[str] = field(default_factory=list)
    checagem: dict = field(default_factory=dict)
    versao_parser: int = VERSAO_PARSER
    # Razão da conta corrente: {seq, data, descricao (sanitizada), valor (com sinal), saldo}.
    # "Saldo Anterior" não é lançamento — vira saldo_inicial_conta.
    saldo_inicial_conta: Optional[float] = None
    lancamentos_conta: list[dict] = field(default_factory=list)
    # Um item por lote do Detalhamento de renda fixa, com a data de aquisição: é como uma
    # compra de Tesouro feita no mês aparece no extrato.
    lotes_rf: list[dict] = field(default_factory=list)
    # v3 — linhas de total de cada bloco, por aba ({aba: [{bloco, rotulo, valores,
    # soma_linhas}]}). Eram descartadas; é delas que o validador confere V2/V3/V5/V6.
    totais_abas: dict = field(default_factory=dict)
    # v3 — fatos atípicos tipados ({tipo, aba, detalhe}): código decorado, quantidade
    # ilegível, fundo encerrado... Nunca param o parse; o validador decide a gravidade.
    avisos_parser: list[dict] = field(default_factory=list)

    def to_dict(self) -> dict:
        d = asdict(self)
        d["posicoes"] = [p.to_dict() for p in self.posicoes]
        return d


# ---------------------------------------------------------------------------
# Identidade do papel entre extratos
# ---------------------------------------------------------------------------
#
# O `nome` do extrato não é identidade: o BTG o reescreve de um mês para o outro
# ('ITAUUNIBANCOPN N1' → 'ITAUUNIBANCOPN EJ N1' quando o papel vai ex-juros). Casar
# posição por nome faz o upsert criar uma linha nova e abandonar a antiga — o ativo
# "renasce" e perde histórico. A chave abaixo é construída do que não muda:
#
#   RV       B3:BBAS3                          ticker da B3
#   TESOURO  TD:LFT:2031-03-01                 sigla do título + vencimento
#   RF       RF:BANCO-XP:CDB:2027-06-15        emissor + sigla + vencimento
#   CAIXA    CAIXA:BTG                         conta corrente da corretora
#
# Renda fixa privada leva o emissor porque dois CDBs de bancos diferentes podem ter
# a mesma sigla e o mesmo vencimento — são papéis distintos, com risco de crédito
# distinto, e não podem colidir numa posição só.

_SEM_VENCIMENTO = "sem-vencimento"


def _chave_rv(ticker: str) -> str:
    return f"B3:{ticker.upper()}"


def _chave_titulo(eh_tesouro: bool, emissor: str, sigla: str, vencimento: Optional[str]) -> str:
    venc = vencimento or _SEM_VENCIMENTO
    if eh_tesouro:
        return f"TD:{_slug(sigla)}:{venc}"
    return f"RF:{_slug(emissor)}:{_slug(sigla)}:{venc}"


# ---------------------------------------------------------------------------
# Privacidade — descrições de texto livre e nome do arquivo
# ---------------------------------------------------------------------------
#
# O razão da conta corrente é a única parte do extrato com texto vindo de fora: um PIX
# recebido traz o nome (às vezes o CPF) de quem pagou; uma TED traz banco, agência e conta.
# O payload vai para o banco (ExtratoImportado) e a descrição vai para a tela, então o que
# sai do parser já sai limpo. Duas regras:
#
#   - linha de transferência (PIX, TED, DOC, TRANSF...): fica só o que está num vocabulário
#     fixo ('PIX RECEBIDO', 'TED ENVIADA', 'TRANSF ENTRE CONTAS'). Nome de pessoa não se
#     reconhece por regex; com lista de permissão, ele simplesmente não passa.
#   - demais linhas: CPF, CNPJ, conta com dígito e sequências longas de dígitos viram '***'.
#     Tickers (ITUB4) e datas (15/05/2029) ficam — é por eles que um provento ou um resgate
#     se casa com o papel.
#
# O nome do XLSX baixado do BTG é o número da conta: mascarar_nome_arquivo vale para ele.

_RE_TRANSFERENCIA = re.compile(r"\b(pix|ted|doc|transf\w*)\b")

_DIRECOES = frozenset({
    "recebido", "recebida", "recebimento", "enviado", "enviada", "envio", "devolvido",
    "devolvida", "devolucao", "estorno", "credito", "debito",
})
_VOCAB_TRANSFERENCIA = _DIRECOES | frozenset({
    "entre", "contas", "conta", "mesma", "titularidade", "custodia", "mercados",
    "investimento", "investimentos", "corrente", "para", "da", "do", "de", "em", "cc", "ci",
    "agendado", "agendada", "programado", "programada", "automatico", "automatica",
})
_SEPARADORES = frozenset({"-", "–", "—", "|", ":", "/", "+", "."})

_MASCARAS = (
    re.compile(r"[\d*]{3}\.[\d*]{3}\.[\d*]{3}-[\d*]{2}"),      # CPF, inteiro ou mascarado
    re.compile(r"\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}"),             # CNPJ
    re.compile(r"(?<!\d)\d{3,}-[\dXx](?![\dA-Za-z])"),           # conta com dígito: 12345-6
    re.compile(r"(?<!\d)\d{5,}(?!\d)"),                          # sequência longa de dígitos
)
_LIMITE_DESCRICAO = 120


def sanitizar_descricao(texto: Any, valor: Optional[float] = None) -> str:
    """Descrição de lançamento sem nome, CPF ou conta de terceiros. Ver o bloco acima."""
    bruto = _txt(texto)
    if not bruto:
        return ""
    if _RE_TRANSFERENCIA.search(_norm(bruto)):
        return _descricao_de_transferencia(bruto, valor)
    limpo = bruto
    for padrao in _MASCARAS:
        limpo = padrao.sub("***", limpo)
    return limpo[:_LIMITE_DESCRICAO]


def _descricao_de_transferencia(bruto: str, valor: Optional[float]) -> str:
    """
    Só o vocabulário fixo sobrevive. Separadores e tokens com dígito (banco, agência) são
    pulados; a primeira palavra desconhecida depois da marca (o nome da contraparte, quase
    sempre) encerra a descrição.
    """
    mantidos: list[str] = []
    viu_marca = False
    for token in bruto.split():
        limpo = token.strip(".,;()")
        n = _norm(limpo)
        if not n or n in _SEPARADORES or any(ch.isdigit() for ch in n):
            continue
        eh_marca = _RE_TRANSFERENCIA.fullmatch(n) is not None
        if eh_marca or n in _VOCAB_TRANSFERENCIA:
            mantidos.append(limpo.upper())
            viu_marca = viu_marca or eh_marca
        elif viu_marca:
            break
    if valor and not any(_norm(t) in _DIRECOES for t in mantidos):
        mantidos.append("RECEBIDO" if valor > 0 else "ENVIADO")
    return " ".join(mantidos)


def mascarar_nome_arquivo(nome: Optional[str]) -> Optional[str]:
    """'001234567.xlsx' → '***.xlsx': o arquivo baixado do BTG leva o número da conta."""
    if not nome:
        return nome
    return re.sub(r"(?<!\d)\d{5,}(?!\d)", "***", nome)


# ---------------------------------------------------------------------------
# Varredura genérica: títulos, cabeçalhos e tabelas
# ---------------------------------------------------------------------------

@dataclass
class _Tabela:
    titulo: str
    header: list[str]         # normalizado
    header_raw: list[str]
    linhas: list[list[Any]] = field(default_factory=list)
    totais: list[list[Any]] = field(default_factory=list)


def _preenchidas(row: list[Any]) -> list[Any]:
    return [c for c in row if c is not None and str(c).strip() != ""]


def _eh_titulo(row: list[Any]) -> bool:
    """'Renda Variável', 'Posição', 'Posição > Ações', 'Posição > Ações | Aluguel'."""
    vals = _preenchidas(row)
    if not vals:
        return False
    if any(isinstance(c, (int, float, datetime, date)) and not isinstance(c, bool) for c in vals):
        return False
    if len(vals) == 1:
        return True
    return ">" in _txt(vals[0])


def _eh_total(row: list[Any]) -> bool:
    vals = _preenchidas(row)
    return bool(vals) and _norm(vals[0]).startswith("total")


def _eh_header(row: list[Any]) -> bool:
    """Linha só de rótulos textuais, com 2+ colunas, que não é título nem total."""
    if _eh_titulo(row) or _eh_total(row):
        return False
    vals = _preenchidas(row)
    if len(vals) < 2:
        return False
    if any(isinstance(c, (int, float, datetime, date)) and not isinstance(c, bool) for c in vals):
        return False
    return all(_num(c) is None for c in vals)


def _tem_numero(row: list[Any]) -> bool:
    """A linha carrega algum número (nativo ou texto '1.234,56')? Datas não contam."""
    return any(
        (isinstance(c, (int, float)) and not isinstance(c, bool)) or _num(c) is not None
        for c in row
    )


def _tabelas(
    ws, aba: str = "", ignoradas: Optional[list] = None, avisos: Optional[list] = None
) -> list[_Tabela]:
    """
    Quebra a planilha em tabelas (título + cabeçalho + linhas + linhas de total).

    v2 — nada se perde em silêncio:
      - linha fora de tabela com célula numérica vai para `ignoradas` com o contexto
        [aba/título] (pode ser dado que o layout novo deslocou);
      - linha fora de tabela só de texto vira um contador em `avisos` (decorativa, mas
        contada);
      - linha com cara de cabeçalho NO MEIO de uma tabela aberta é DADO: uma linha de
        dados cheia de '-' (texto sem número, célula a célula) é indistinguível de um
        cabeçalho. Cabeçalho só abre tabela nova quando não há tabela aberta ou quando a
        atual já fechou com linha(s) de total — nos extratos reais todo bloco novo vem
        depois de um título ou de uma linha em branco, então nada muda para eles.
    """
    tabelas: list[_Tabela] = []
    titulo = ""
    atual: Optional[_Tabela] = None
    so_texto = 0

    for raw in ws.iter_rows(values_only=True):
        row = list(raw)
        if not _preenchidas(row):
            atual = None
            continue
        if _eh_titulo(row):
            titulo = " | ".join(_txt(c) for c in _preenchidas(row))
            atual = None
            continue
        if _eh_header(row) and (atual is None or atual.totais):
            atual = _Tabela(
                titulo=titulo,
                header=[_norm(c) for c in row],
                header_raw=[_txt(c) for c in row],
            )
            tabelas.append(atual)
            continue
        if atual is None:
            if _tem_numero(row):
                if ignoradas is not None:
                    contexto = "/".join(p for p in (aba, titulo) if p)
                    ignoradas.append(_resumo_linha(contexto, row))
            else:
                so_texto += 1
            continue
        if _eh_total(row):
            atual.totais.append(row)
        else:
            atual.linhas.append(row)

    if so_texto:
        _avisar(
            avisos, "linhas_texto_fora_de_tabela", aba or _txt(ws.title),
            f"{so_texto} linha(s) só de texto fora de tabela descartada(s)",
        )
    return tabelas


def _dict_total(bloco: str, row: list[Any], header: list[str]) -> dict:
    """Linha de total → {bloco, rotulo, valores: {coluna normalizada: float|None}}."""
    vals = _preenchidas(row)
    valores: dict[str, Optional[float]] = {}
    for i, h in enumerate(header):
        if h:
            valores[h] = _num(_cel(row, i))
    return {"bloco": bloco, "rotulo": _txt(vals[0]) if vals else "", "valores": valores}


def _totais_da_tabela(tab: _Tabela) -> list[dict]:
    """
    Entradas de `totais_abas` para um bloco: cada linha de total + `soma_linhas`, a soma
    das linhas de dados nas mesmas colunas. É a dupla que o validador compara (Σ linhas =
    Total do bloco; Σ totais = Sumário) sem reparsear o extrato.
    """
    soma: dict[str, float] = {}
    for i, h in enumerate(tab.header):
        if not h:
            continue
        numeros = [n for n in (_num(_cel(r, i)) for r in tab.linhas) if n is not None]
        if numeros:
            soma[h] = round(sum(numeros), 2)
    saida: list[dict] = []
    for tot in tab.totais:
        entrada = _dict_total(tab.titulo, tot, tab.header)
        entrada["soma_linhas"] = soma
        saida.append(entrada)
    return saida


def _idx(tab: _Tabela, *alvos: str) -> Optional[int]:
    """Índice da coluna pelo rótulo: casamento exato primeiro, depois por prefixo."""
    normalizados = [_norm(a) for a in alvos]
    for alvo in normalizados:
        for i, h in enumerate(tab.header):
            if h == alvo:
                return i
    for alvo in normalizados:
        for i, h in enumerate(tab.header):
            if h and (h.startswith(alvo) or alvo in h):
                return i
    return None


def _cel(row: list[Any], i: Optional[int]) -> Any:
    if i is None or i >= len(row):
        return None
    return row[i]


def _aba(wb, nome: str):
    """Aba pelo nome normalizado — o título tem acento ('Valores em Trânsito')."""
    alvo = _norm(nome)
    for ws in wb.worksheets:
        if _norm(ws.title) == alvo:
            return ws
    for ws in wb.worksheets:
        if alvo in _norm(ws.title):
            return ws
    return None


def _resumo_linha(titulo: str, row: list[Any]) -> str:
    conteudo = " | ".join(_txt(c) for c in _preenchidas(row))
    return f"[{titulo}] {conteudo}"[:300]


# ---------------------------------------------------------------------------
# Capa — data de referência
# ---------------------------------------------------------------------------

def _parse_capa(wb) -> tuple[Optional[str], dict]:
    ws = _aba(wb, "Capa")
    if ws is None:
        raise ExtratoParseError("Aba 'Capa' não encontrada — o arquivo não parece um extrato do BTG.")

    data_ref: Optional[str] = None
    emitido_em: Optional[str] = None
    periodo_inicio: Optional[str] = None

    for row in ws.iter_rows(values_only=True):
        for cell in row:
            texto = _txt(cell)
            if not texto:
                continue
            n = _norm(texto)
            if n.startswith("periodo de"):
                datas = re.findall(r"\d{2}/\d{2}/\d{2,4}", texto)
                if len(datas) >= 2:
                    periodo_inicio, data_ref = _data(datas[0]), _data(datas[1])
                elif datas:
                    data_ref = _data(datas[0])
            elif n.startswith("emitido em"):
                emitido_em = _data(texto)

    if not data_ref:
        raise ExtratoParseError(
            "Não encontrei o período do extrato na aba 'Capa' "
            "(linha 'Período de DD/MM/AA a DD/MM/AA'). Sem data de referência não importo."
        )
    return data_ref, {"periodo_inicio": periodo_inicio, "emitido_em": emitido_em}


# ---------------------------------------------------------------------------
# Sumario — totais do mês atual e do anterior (base do checksum)
# ---------------------------------------------------------------------------

def _parse_sumario(wb) -> dict:
    ws = _aba(wb, "Sumario")
    if ws is None:
        return {}

    tab = next((t for t in _tabelas(ws) if _idx(t, "mercados") is not None), None)
    if tab is None:
        return {}

    i_mercado = _idx(tab, "mercados")

    # Rótulos são 'Saldo Bruto R$ 30/06/26' — as datas mudam todo mês, então a coluna é
    # identificada por bruto/líquido + data extraída do próprio rótulo.
    colunas: list[tuple[int, str, Optional[str]]] = []
    for i, raw in enumerate(tab.header_raw):
        n = _norm(raw)
        if not n.startswith("saldo"):
            continue
        tipo = "bruto" if "bruto" in n else "liquido"
        colunas.append((i, tipo, _data(raw)))

    if not colunas:
        return {}

    datas = sorted({d for _, _, d in colunas if d})
    data_atual = datas[-1] if datas else None
    data_anterior = datas[0] if len(datas) > 1 else None

    def _periodo(data_alvo: Optional[str]) -> dict:
        alvo = [(i, tipo) for i, tipo, d in colunas if d == data_alvo]
        if not alvo and data_alvo is None:
            alvo = [(i, tipo) for i, tipo, _ in colunas]
        mercados: dict[str, dict] = {}
        for row in tab.linhas:
            nome = _txt(_cel(row, i_mercado))
            if not nome:
                continue
            mercados[nome] = {tipo: _num(_cel(row, i)) for i, tipo in alvo}
        totais: dict[str, Optional[float]] = {}
        if tab.totais:
            totais = {tipo: _num(_cel(tab.totais[0], i)) for i, tipo in alvo}
        return {"data": data_alvo, "mercados": mercados, "total": totais}

    sumario = {"atual": _periodo(data_atual)}
    if data_anterior:
        sumario["anterior"] = _periodo(data_anterior)
    return sumario


def _norm_mercado(nome: Any) -> str:
    """'Renda Variável*' → 'renda variavel': período aberto decora até o NOME do mercado."""
    return _norm(nome).replace("*", "").strip()


def _sumario_mercado(sumario: dict, chave: str, periodo: str = "atual") -> Optional[float]:
    """Saldo bruto de um mercado do Sumário ('atual' ou 'anterior'). '-' no extrato → None."""
    bloco = sumario.get(periodo) or {}
    for nome, valores in (bloco.get("mercados") or {}).items():
        if _norm_mercado(nome) == _norm_mercado(chave):
            return valores.get("bruto")
    return None


# ---------------------------------------------------------------------------
# Renda Variável
# ---------------------------------------------------------------------------

def _parse_renda_variavel(
    wb, as_of: str, ignoradas: list[str], avisos: Optional[list] = None
) -> tuple[list[PosicaoParsed], list[dict], list[dict], list[dict], list[dict], list[dict]]:
    """(posições, proventos, movimentações, aluguel, conferência por bloco, totais da aba)."""
    ws = _aba(wb, "Renda Variavel")
    posicoes: list[PosicaoParsed] = []
    proventos: list[dict] = []
    movimentacoes: list[dict] = []
    aluguel: list[dict] = []
    blocos: list[dict] = []
    totais: list[dict] = []
    if ws is None:
        return posicoes, proventos, movimentacoes, aluguel, blocos, totais

    for tab in _tabelas(ws, "Renda Variavel", ignoradas, avisos):
        t = _norm(tab.titulo)
        eh_aluguel = "aluguel" in t

        if t.startswith("posicao >") and eh_aluguel:
            aluguel.extend(_parse_aluguel(tab, avisos))
        elif t.startswith("posicao >"):
            classe = _classe_rv(t)
            if classe is None:
                ignoradas.append(f"[{tab.titulo}] bloco de posição não reconhecido — {len(tab.linhas)} linha(s)")
                continue
            posicoes.extend(_parse_rv(tab, classe, as_of, ignoradas, avisos))
        elif t.startswith("movimentacao >") and not eh_aluguel:
            p, m = _parse_movimentacao(tab, avisos)
            proventos.extend(p)
            movimentacoes.extend(m)
            blocos.append(_checagem_bloco_rv(tab, p, m))
        # 'Movimentação > Ações | Aluguel' é redundante com a posição de aluguel: ignorado.
        totais.extend(_totais_da_tabela(tab))

    return posicoes, proventos, movimentacoes, aluguel, blocos, totais


def _classe_rv(titulo_norm: str) -> Optional[str]:
    depois = titulo_norm.split(">", 1)[1].strip() if ">" in titulo_norm else titulo_norm
    for marca, classe in _CLASSE_RV.items():
        if marca in depois:
            return classe
    return None


def _parse_rv(
    tab: _Tabela, classe: str, as_of: str, ignoradas: list[str], avisos: Optional[list] = None
) -> list[PosicaoParsed]:
    i_cod = _idx(tab, "codigo")
    i_nome = _idx(tab, "acao", "ativo")
    i_qtde = _idx(tab, "qtde", "quantidade")
    i_fech = _idx(tab, "preco fechamento")
    i_medio = _idx(tab, "preco medio")
    i_saldo = _idx(tab, "saldo bruto")

    out: list[PosicaoParsed] = []
    for row in tab.linhas:
        ticker, decorado = _limpar_codigo(_cel(row, i_cod))
        saldo = _num(_cel(row, i_saldo))
        if not ticker or saldo is None:
            ignoradas.append(_resumo_linha(tab.titulo, row))
            continue
        if decorado:
            _avisar(avisos, "ticker_decorado", "Renda Variavel",
                    f"{ticker}: código veio decorado ('*') — extrato de período aberto")
        if not _RE_TICKER_B3.match(ticker):
            # Formato estranho nunca descarta a posição: o dinheiro existe mesmo que o
            # código não case com o padrão da B3 (subscrição, fracionário novo...).
            _avisar(avisos, "ticker_fora_do_padrao_b3", "Renda Variavel",
                    f"{ticker}: código fora do padrão da B3 — posição mantida")
        preco = _num(_cel(row, i_fech))
        qtde = _num(_cel(row, i_qtde))
        if qtde is None:
            qtde, tipo = _derivar_quantidade(saldo, preco)
            _avisar(avisos, tipo, "Renda Variavel",
                    f"{ticker}: quantidade ilegível — "
                    + ("derivada de saldo/preço" if tipo == "quantidade_derivada" else "ficou 0.0"))
        out.append(PosicaoParsed(
            ticker=ticker,
            nome=_txt(_cel(row, i_nome)) or ticker,
            classe=classe,
            quantidade=qtde,
            preco_medio=_num(_cel(row, i_medio)),
            valor_mercado=saldo,
            preco_fechamento=preco,
            as_of=as_of,
            chave_externa=_chave_rv(ticker),
        ))
    return out


def _parse_aluguel(tab: _Tabela, avisos: Optional[list] = None) -> list[dict]:
    """
    Ações doadas em aluguel. NÃO é posição: o papel já está contado na posição de Ações
    (TAEE11 aparece nos dois blocos). Somar duplicaria o ativo.
    """
    i_cod = _idx(tab, "codigo")
    i_qtde = _idx(tab, "qtde", "quantidade")
    i_pos = _idx(tab, "posicao")
    i_valor = _idx(tab, "valor contratado")
    i_venc = _idx(tab, "data vencimento")
    i_taxa = _idx(tab, "taxa ano", "taxa")
    i_result = _idx(tab, "result")

    out: list[dict] = []
    for row in tab.linhas:
        cod, decorado = _limpar_codigo(_cel(row, i_cod))
        if not cod:
            continue
        if decorado:
            _avisar(avisos, "ticker_decorado", "Renda Variavel",
                    f"{cod}: código veio decorado ('*') — extrato de período aberto")
        out.append({
            "ticker": cod,
            "quantidade": _num(_cel(row, i_qtde)),
            "posicao": _txt(_cel(row, i_pos)),
            "valor_contratado": _num(_cel(row, i_valor)),
            "vencimento": _data(_cel(row, i_venc)),
            "taxa_ano_pct": _num(_cel(row, i_taxa)),
            "resultado_acumulado_liquido": _num(_cel(row, i_result)),
        })
    return out


def _operacao_rv(transacao_norm: str) -> str:
    """COMPRA | VENDA | PROVENTO | OUTRO (bonificação, subscrição, transferência...)."""
    if transacao_norm.startswith("compra"):
        return "COMPRA"
    if transacao_norm.startswith("venda"):
        return "VENDA"
    if any(m in transacao_norm for m in _PROVENTO_MARCAS):
        return "PROVENTO"
    return "OUTRO"


def _parse_movimentacao(tab: _Tabela, avisos: Optional[list] = None) -> tuple[list[dict], list[dict]]:
    i_data = _idx(tab, "data")
    i_trans = _idx(tab, "transacao")
    i_cod = _idx(tab, "codigo")
    i_qtde = _idx(tab, "qtde", "quantidade")
    i_preco = _idx(tab, "preco r$", "preco")
    i_bruto = _idx(tab, "valor bruto")
    i_corretagem = _idx(tab, "corretagem")
    i_liq = _idx(tab, "valor liquido")
    # A classe sai do título do bloco ('Movimentação > Fundos Listados' → FII): é o que
    # permite somar a renda e as compras por classe sem depender da posição do mês.
    classe = _classe_rv(_norm(tab.titulo))

    proventos: list[dict] = []
    outras: list[dict] = []
    for row in tab.linhas:
        transacao = _txt(_cel(row, i_trans))
        if not transacao:
            continue
        n = _norm(transacao)
        operacao = _operacao_rv(n)
        ticker, decorado = _limpar_codigo(_cel(row, i_cod))
        if decorado:
            _avisar(avisos, "ticker_decorado", "Renda Variavel",
                    f"{ticker}: código veio decorado ('*') — extrato de período aberto")
        item = {
            "data": _data(_cel(row, i_data)),
            "transacao": transacao,
            "ticker": ticker or None,
            "quantidade": _num(_cel(row, i_qtde)),
            "valor_bruto": _num(_cel(row, i_bruto)),
            "valor_liquido": _num(_cel(row, i_liq)),
            "classe": classe,
            "operacao": operacao,
            "preco": _num(_cel(row, i_preco)),
            "corretagem": _num(_cel(row, i_corretagem)),
            "amortizacao": "amortizacao" in n,
        }
        (proventos if operacao == "PROVENTO" else outras).append(item)
    return proventos, outras


def _checagem_bloco_rv(tab: _Tabela, proventos: list[dict], outras: list[dict]) -> dict:
    """
    Confere o que foi lido contra as linhas 'Total de Compras/Vendas/Proventos' do bloco.
    Nos extratos reais a linha de total traz mais de um número (bruto, corretagem,
    líquido), então a conferência é tolerante: basta a soma lida bater com algum deles.
    """
    def _liquido(item: dict) -> float:
        valor = item.get("valor_liquido")
        if valor is None:
            valor = item.get("valor_bruto")
        return abs(valor or 0.0)

    lido = {
        "compras": round(sum(_liquido(i) for i in outras if i["operacao"] == "COMPRA"), 2),
        "vendas": round(sum(_liquido(i) for i in outras if i["operacao"] == "VENDA"), 2),
        "proventos": round(sum(_liquido(i) for i in proventos), 2),
    }
    extrato: dict[str, list[float]] = {}
    for tot in tab.totais:
        preenchidas = _preenchidas(tot)
        rotulo = _norm(preenchidas[0]) if preenchidas else ""
        numeros = [abs(n) for n in (_num(c) for c in tot) if n is not None]
        for chave in lido:
            if chave[:-1] in rotulo:          # 'compra' em 'total de compras'
                extrato[chave] = numeros

    ok = True
    for chave, soma in lido.items():
        numeros = extrato.get(chave)
        if numeros is None:                   # bloco sem essa linha de total
            continue
        if not numeros:                       # total '-' → nada no mês
            ok = ok and soma <= TOLERANCIA_CONTA
        else:
            ok = ok and any(abs(soma - n) <= TOLERANCIA_CONTA for n in numeros)
    return {"bloco": tab.titulo, "ok": ok, "lido": lido, "extrato": extrato}


# ---------------------------------------------------------------------------
# Renda Fixa (Tesouro Direto e, se houver, RF privada)
# ---------------------------------------------------------------------------

def _chave_rf(sigla: str, vencimento: Optional[str]) -> tuple[str, str]:
    """Chave de casamento posição ↔ lote de aquisição, dentro de um mesmo extrato."""
    return (sigla.upper(), vencimento or "")


def _parse_renda_fixa(
    wb, as_of: str, ignoradas: list[str], avisos: Optional[list] = None
) -> tuple[list[PosicaoParsed], list[dict], list[dict]]:
    """(posições, lotes, totais da aba) — os lotes, com data de aquisição, vão para `lotes_rf`."""
    ws = _aba(wb, "Renda Fixa")
    if ws is None:
        return [], [], []

    # Posições são agregadas pela identidade do papel (chave_externa); os lotes de
    # aquisição casam por (sigla, vencimento), que é o que o Detalhamento traz.
    posicoes: dict[str, PosicaoParsed] = {}
    chaves_lote: dict[str, tuple[tuple[str, str], str]] = {}
    lotes: dict[tuple[str, str], dict[str, dict]] = {}
    lotes_saida: list[dict] = []
    totais: list[dict] = []

    for tab in _tabelas(ws, "Renda Fixa", ignoradas, avisos):
        t = _norm(tab.titulo)
        if t.startswith("posicao consolidada"):
            continue  # conferência por emissor — nada a extrair
        if t.startswith("posicao >"):
            _acumula_posicao_rf(tab, as_of, posicoes, chaves_lote, ignoradas, avisos)
        elif t.startswith("detalhamento >"):
            _acumula_lotes_rf(tab, lotes, lotes_saida, avisos)
        totais.extend(_totais_da_tabela(tab))

    # Custo real de aquisição: só existe no Detalhamento. Sem lote casado, preco_medio
    # fica None — nunca o preço atual disfarçado de custo (era o erro do parser de PDF).
    disputados = Counter(cl for cl, _ in chaves_lote.values())
    for chave, pos in posicoes.items():
        chave_lote, emissor_slug = chaves_lote[chave]
        lote = _lote_do_papel(
            lotes.get(chave_lote), emissor_slug, disputado=disputados[chave_lote] > 1
        )
        if lote and lote["quantidade"]:
            # Média ponderada pelo preço de compra declarado; o valor de compra (já
            # arredondado em centavos pelo BTG) é o fallback.
            base = lote["preco_ponderado"] or lote["valor_compra"]
            pos.preco_medio = round(base / lote["quantidade"], 2)
            pos.custo_total = round(lote["valor_compra"], 2)
        else:
            logger.info("renda fixa: sem lote de aquisição para %s — preco_medio ficará nulo", chave)

    return list(posicoes.values()), lotes_saida, totais


def _lote_do_papel(
    por_emissor: Optional[dict[str, dict]], emissor_slug: str, disputado: bool
) -> Optional[dict]:
    """
    Escolhe o lote de aquisição do papel. O Detalhamento não traz coluna de emissor —
    ele vem no título do bloco —, então com um único bloco e um único papel disputando
    a chave o casamento é direto.

    `disputado` = mais de um papel (emissores diferentes, mesma sigla e vencimento)
    aponta para esta chave de lote. Aí só o emissor exato serve: aplicar o lote errado
    inventaria um custo médio, e custo inventado é pior que custo ausente.
    """
    if not por_emissor:
        return None
    if len(por_emissor) == 1 and not disputado:
        return next(iter(por_emissor.values()))
    return por_emissor.get(emissor_slug)


def _acumula_posicao_rf(
    tab: _Tabela,
    as_of: str,
    posicoes: dict[str, PosicaoParsed],
    chaves_lote: dict[str, tuple[tuple[str, str], str]],
    ignoradas: list[str],
    avisos: Optional[list] = None,
) -> None:
    i_emissor = _idx(tab, "emissor")
    i_ativo = _idx(tab, "ativo")
    i_venc = _idx(tab, "vencimento")
    i_taxa = _idx(tab, "taxa media ponderada", "taxa")
    i_qtde = _idx(tab, "quantidade", "qtde")
    i_preco = _idx(tab, "preco r$", "preco")
    i_saldo = _idx(tab, "saldo bruto")

    # 'Posição > TESOURO DIRETO - LFT' — usado quando a linha não traz a coluna Ativo
    sigla_titulo = tab.titulo.split("-")[-1].strip().upper() if "-" in tab.titulo else ""
    eh_tesouro_titulo = "tesouro direto" in _norm(tab.titulo)

    for row in tab.linhas:
        saldo = _num(_cel(row, i_saldo))
        # Limpeza ANTES do lookup em _TESOURO_NOME: 'LFT*' (período aberto) tem de achar
        # 'Tesouro Selic', e a chave_externa tem de sair sem decoração.
        sigla, decorado = _limpar_codigo(_txt(_cel(row, i_ativo)) or sigla_titulo)
        celula_venc = _cel(row, i_venc)
        vencimento = _data(celula_venc)
        if saldo is None or not sigla:
            ignoradas.append(_resumo_linha(tab.titulo, row))
            continue
        if decorado:
            _avisar(avisos, "ticker_decorado", "Renda Fixa",
                    f"{sigla}: sigla veio decorada ('*') — extrato de período aberto")
        if vencimento is None:
            # Sem vencimento a chave cai em 'sem-vencimento' — legível ou não, é fato
            # que o validador precisa conhecer (dois títulos ali colidem).
            texto_venc = _txt(celula_venc)
            tipo_venc = "vencimento_ilegivel" if texto_venc and texto_venc != "-" else "vencimento_ausente"
            _avisar(avisos, tipo_venc, "Renda Fixa",
                    f"{sigla}: vencimento {'ilegível' if tipo_venc == 'vencimento_ilegivel' else 'ausente'} "
                    "— papel agregado em 'sem-vencimento'")

        emissor = _txt(_cel(row, i_emissor))
        eh_tesouro = eh_tesouro_titulo or "BACEN" in emissor.upper()
        ano = _ano(vencimento)
        preco = _num(_cel(row, i_preco))
        qtde = _num(_cel(row, i_qtde))
        if qtde is None:
            qtde, tipo = _derivar_quantidade(saldo, preco)
            _avisar(avisos, tipo, "Renda Fixa",
                    f"{sigla}: quantidade ilegível — "
                    + ("derivada de saldo/preço" if tipo == "quantidade_derivada" else "ficou 0.0"))
        taxa = _txt(_cel(row, i_taxa)) or None

        emissor_curto = emissor.split(" - ")[0].strip()[:40]
        if eh_tesouro:
            nome = f"{_TESOURO_NOME.get(sigla, sigla)} {ano}"
            ticker: Optional[str] = nome
            classe = "TESOURO"
        else:
            nome = f"{sigla} {emissor_curto or 'emissor não identificado'} {ano}".strip()
            ticker = None
            classe = "RF"

        chave = _chave_titulo(eh_tesouro, emissor_curto, sigla, vencimento)
        existente = posicoes.get(chave)
        if existente:
            if vencimento is None:
                # Dois papéis fundidos por falta de vencimento legível: pode ser o mesmo
                # título em duas emissões — ou dois títulos diferentes. O validador
                # promove este aviso a erro (a posição agregada seria mentira).
                _avisar(avisos, "colisao_sem_vencimento", "Renda Fixa",
                        f"{chave}: mais de um papel agregado sem vencimento legível")
            # Mesmo papel em duas linhas (emissões diferentes do mesmo título): agrega.
            existente.quantidade += qtde
            existente.valor_mercado += saldo
        else:
            chaves_lote[chave] = (_chave_rf(sigla, vencimento), _slug(emissor_curto))
            posicoes[chave] = PosicaoParsed(
                ticker=ticker,
                nome=nome,
                classe=classe,
                quantidade=qtde,
                preco_medio=None,
                valor_mercado=saldo,
                preco_fechamento=preco,
                as_of=as_of,
                vencimento=vencimento,
                taxa_contratada=taxa,
                chave_externa=chave,
            )


def _acumula_lotes_rf(
    tab: _Tabela,
    lotes: dict[tuple[str, str], dict[str, dict]],
    lotes_saida: list[dict],
    avisos: Optional[list] = None,
) -> None:
    i_ativo = _idx(tab, "ativo")
    i_venc = _idx(tab, "vencimento")
    i_qtde = _idx(tab, "quantidade", "qtde")
    i_valor = _idx(tab, "valor compra")
    i_preco = _idx(tab, "preco compra")
    # Casamento exato de propósito: 'data' casaria com 'Data inicial de liquidez' e
    # 'preco' com 'Preço Compra R$'.
    i_aquisicao = _idx(tab, "aquisicao")
    i_taxa = _idx(tab, "taxa compra")
    i_preco_atual = _idx(tab, "preco r$")
    i_saldo = _idx(tab, "saldo bruto")

    # 'Detalhamento > TESOURO DIRETO - LFT | BACEN-BANCO CENTRAL DO BRASIL - RJ':
    # a sigla vem antes do '|', o emissor depois — as linhas não repetem o emissor.
    cabecalho, _, emissor_titulo = tab.titulo.partition("|")
    sigla_titulo = cabecalho.split("-")[-1].strip().upper() if "-" in cabecalho else ""
    emissor_curto = emissor_titulo.split(" - ")[0].strip()[:40]
    emissor_slug = _slug(emissor_curto)
    # Mesma regra de _acumula_posicao_rf: a chave do lote tem de ser a da posição.
    eh_tesouro = "tesouro direto" in _norm(cabecalho) or "BACEN" in emissor_curto.upper()

    for row in tab.linhas:
        # Mesma limpeza da posição: a chave do lote tem de casar com a chave do papel
        # mesmo quando só um dos dois lados veio decorado ('LFT*' na posição, 'LFT' aqui).
        sigla, decorado = _limpar_codigo(_txt(_cel(row, i_ativo)) or sigla_titulo)
        qtde = _num(_cel(row, i_qtde))
        valor = _num(_cel(row, i_valor))
        if not sigla or qtde is None or valor is None:
            continue
        if decorado:
            _avisar(avisos, "ticker_decorado", "Renda Fixa",
                    f"{sigla}: sigla veio decorada ('*') — extrato de período aberto")
        preco = _num(_cel(row, i_preco))
        vencimento = _data(_cel(row, i_venc))
        chave = _chave_rf(sigla, vencimento)
        acc = lotes.setdefault(chave, {}).setdefault(
            emissor_slug, {"quantidade": 0.0, "valor_compra": 0.0, "preco_ponderado": 0.0, "lotes": 0}
        )
        acc["quantidade"] += qtde
        acc["valor_compra"] += valor
        acc["preco_ponderado"] += qtde * preco if preco is not None else valor
        acc["lotes"] += 1

        lotes_saida.append({
            "chave_externa": _chave_titulo(eh_tesouro, emissor_curto, sigla, vencimento),
            "sigla": sigla,
            "vencimento": vencimento,
            "aquisicao": _data(_cel(row, i_aquisicao)),
            "quantidade": qtde,
            "preco_compra": preco,
            "valor_compra": valor,
            "taxa_compra": _txt(_cel(row, i_taxa)) or None,
            "preco_atual": _num(_cel(row, i_preco_atual)),
            "saldo_bruto": _num(_cel(row, i_saldo)),
        })


# ---------------------------------------------------------------------------
# Fundos (aba "Fundos") e CriptoAtivos
# ---------------------------------------------------------------------------
#
# Chaves: FUNDO:{slug do CNPJ da classe, quando presente no nome, senão slug do nome} e
# CRIPTO:{símbolo entre parênteses no Ativo}. O CNPJ identifica a classe do fundo mesmo
# que o BTG reescreva o nome; o símbolo ('XBT BITCOIN (XBT)' → XBT) é o que não muda.

_RE_CNPJ = re.compile(r"\d{2}\.?\d{3}\.?\d{3}/\d{4}-?\d{0,2}")
_RE_SIMBOLO_CRIPTO = re.compile(r"\(([A-Za-z0-9]+)\)\s*$")


def _nome_e_chave_fundo(nome_bruto: str) -> tuple[str, str]:
    """('Fundo X - Classe CNPJ: 11.222.333/0001-44') → ('Fundo X', 'FUNDO:11-222-333-0001-44')."""
    nome = re.sub(r"\s*-?\s*classe\s+cnpj:.*$", "", nome_bruto, flags=re.IGNORECASE).strip()
    nome = nome or nome_bruto or "Fundo sem nome"
    m = _RE_CNPJ.search(nome_bruto)
    return nome, f"FUNDO:{_slug(m.group(0)) if m else _slug(nome)}"


def _parse_fundos(
    wb, as_of: str, ignoradas: list[str], avisos: Optional[list] = None
) -> tuple[list[PosicaoParsed], list[dict], list[dict], list[dict]]:
    """
    (posições FUNDO, proventos, movimentações, totais da aba).

    A aba tem um leitor próprio em vez de _tabelas por causa do layout da seção
    'Posição > Portfólio de fundos': cada fundo ocupa DUAS linhas — a primeira só com o
    nome ('... - Classe CNPJ: NN.NNN.NNN/NNNN-NN') numa célula, a segunda com os dados.
    A linha de nome tem uma célula única, que _eh_titulo tomaria por título de seção, e
    a linha de dados viraria órfã.

    Fundo encerrado no período: as células atuais vêm '-' e só o Saldo Líquido do período
    anterior (coluna com data no rótulo) vem preenchido. NÃO vira posição — inventar uma
    posição de R$ 0,00 sujaria a carteira — mas vira aviso, para a ausência não passar
    por perda de parse.
    """
    posicoes: list[PosicaoParsed] = []
    proventos: list[dict] = []
    movs: list[dict] = []
    totais: list[dict] = []
    ws = _aba(wb, "Fundos")
    if ws is None:
        return posicoes, proventos, movs, totais

    _SECOES = {"fundos", "posicao", "posicoes", "movimentacao", "movimentacoes"}
    modo = ""                       # "posicao" | "mov" | ""
    header: list[str] = []
    header_raw: list[str] = []
    bloco = ""                      # título do bloco corrente (para totais_abas)
    fundo_mov = ""                  # nome do fundo do bloco de movimentação corrente
    nome_pendente = ""              # linha 1 (nome) do registro de posição corrente
    soma_linhas: dict[str, float] = {}
    chaves: dict[str, str] = {}     # nome normalizado → chave (liga movimentação a fundo)

    def _acumular_soma(row: list[Any]) -> None:
        for i, h in enumerate(header):
            if not h:
                continue
            n = _num(_cel(row, i))
            if n is not None:
                soma_linhas[h] = round(soma_linhas.get(h, 0.0) + n, 2)

    for raw in ws.iter_rows(values_only=True):
        row = list(raw)
        vals = _preenchidas(row)
        if not vals:
            nome_pendente = ""
            continue
        n0 = _norm(vals[0])
        if n0.startswith("posicao >"):
            modo, header, nome_pendente, bloco = "posicao", [], "", _txt(vals[0])
            soma_linhas = {}
            continue
        if n0.startswith("movimentacao >"):
            modo, header, bloco = "mov", [], _txt(vals[0])
            fundo_mov = bloco.split(">", 1)[1].strip()
            soma_linhas = {}
            continue
        if n0 in _SECOES:
            modo = ""
            continue
        if not modo:
            continue
        if not header:
            if _eh_header(row):
                header = [_norm(c) for c in row]
                header_raw = [_txt(c) for c in row]
            continue
        if _eh_total(row):
            entrada = _dict_total(bloco, row, header)
            entrada["soma_linhas"] = dict(soma_linhas)
            totais.append(entrada)
            if modo == "posicao":
                modo = ""           # 'Total em fundos' fecha a seção de posições
            continue
        if modo == "posicao":
            if len(vals) == 1 and _num(vals[0]) is None and not isinstance(vals[0], (datetime, date)):
                nome_pendente = _txt(vals[0])
                continue
            _acumular_soma(row)
            _fundo_posicao(row, header, header_raw, nome_pendente, as_of,
                           posicoes, chaves, ignoradas, avisos)
            nome_pendente = ""
        else:
            _acumular_soma(row)
            _fundo_movimentacao(row, header, fundo_mov, chaves, proventos, movs,
                                ignoradas, avisos)

    return posicoes, proventos, movs, totais


def _fundo_posicao(
    row: list[Any],
    header: list[str],
    header_raw: list[str],
    nome_bruto: str,
    as_of: str,
    posicoes: list[PosicaoParsed],
    chaves: dict[str, str],
    ignoradas: list[str],
    avisos: Optional[list],
) -> None:
    def _col(alvo: str, exato: bool = False, com_data: Optional[bool] = None) -> Optional[int]:
        for i, h in enumerate(header):
            if (h == alvo if exato else h.startswith(alvo)):
                if com_data is None or bool(_data(header_raw[i])) == com_data:
                    return i
        return None

    i_ref = _col("data referencia")
    i_ant = _col("saldo liquido", com_data=True)     # 'Saldo Líquido R$ dd/mm/aa' = período anterior
    i_qtde = _col("quantidade de cotas")
    if i_qtde is None:
        i_qtde = _col("quantidade")
    i_cot = _col("cotacao")
    i_bruto = _col("saldo bruto")
    i_liq = _col("saldo liquido r$", exato=True)
    if i_liq is None:
        i_liq = _col("saldo liquido", com_data=False)

    qtde = _num(_cel(row, i_qtde))
    cotacao = _num(_cel(row, i_cot))
    bruto = _num(_cel(row, i_bruto))
    liq = _num(_cel(row, i_liq))
    anterior = _num(_cel(row, i_ant))

    nome, chave = _nome_e_chave_fundo(nome_bruto)
    if nome_bruto:
        chaves[_norm(nome)] = chave

    if bruto is None and liq is None and qtde is None:
        if anterior is not None:
            # Encerrado no período: só o saldo do período anterior vem preenchido.
            _avisar(avisos, "fundo_encerrado", "Fundos",
                    f"{nome}: fundo encerrado no período — sem posição atual")
        else:
            ignoradas.append(_resumo_linha(f"Fundos/{nome or 'posição'}", row))
        return

    if not nome_bruto:
        _avisar(avisos, "fundo_sem_nome", "Fundos",
                "linha de dados sem a linha de nome logo acima — chave caiu no genérico")
    valor = bruto if bruto is not None else liq
    if bruto is None:
        _avisar(avisos, "fundo_sem_saldo_bruto", "Fundos",
                f"{nome}: sem Saldo Bruto — usei o Saldo Líquido como valor")
    if qtde is None:
        qtde, tipo = _derivar_quantidade(valor or 0.0, cotacao)
        _avisar(avisos, tipo, "Fundos",
                f"{nome}: quantidade de cotas ilegível — "
                + ("derivada de saldo/cotação" if tipo == "quantidade_derivada" else "ficou 0.0"))

    posicoes.append(PosicaoParsed(
        ticker=None,
        nome=nome,
        classe="FUNDO",
        quantidade=qtde,
        preco_medio=None,
        valor_mercado=valor or 0.0,
        preco_fechamento=cotacao,
        as_of=_data(_cel(row, i_ref)) or as_of,
        chave_externa=chave,
    ))


def _operacao_fundo(transacao_norm: str) -> str:
    """Aplicação ≈ compra e resgate ≈ venda: é assim que o motor de desempenho os lê."""
    if transacao_norm.startswith("aplic"):
        return "COMPRA"
    if transacao_norm.startswith("resgat"):
        return "VENDA"
    return _operacao_rv(transacao_norm)


def _fundo_movimentacao(
    row: list[Any],
    header: list[str],
    fundo: str,
    chaves: dict[str, str],
    proventos: list[dict],
    movs: list[dict],
    ignoradas: list[str],
    avisos: Optional[list],
) -> None:
    def _col(*alvos: str) -> Optional[int]:
        for alvo in alvos:
            for i, h in enumerate(header):
                if h.startswith(alvo):
                    return i
        return None

    transacao = _txt(_cel(row, _col("transacao")))
    if not transacao:
        if _tem_numero(row):
            ignoradas.append(_resumo_linha(f"Fundos/Movimentação {fundo}", row))
        return
    n = _norm(transacao)
    operacao = _operacao_fundo(n)
    # ticker fica None de propósito: quem consome movimentações hoje monta chave B3 a
    # partir dele; o fundo se identifica pela chave_externa (e pelo nome, para a tela).
    item = {
        "data": _data(_cel(row, _col("data"))),
        "transacao": transacao,
        "ticker": None,
        "quantidade": _num(_cel(row, _col("quantidade de cotas", "quantidade", "qtde"))),
        "valor_bruto": _num(_cel(row, _col("valor bruto"))),
        "valor_liquido": _num(_cel(row, _col("valor liquido"))),
        "classe": "FUNDO",
        "operacao": operacao,
        "preco": _num(_cel(row, _col("valor da cota", "preco"))),
        "corretagem": None,
        "amortizacao": "amortizacao" in n,
        "nome": fundo,
        "chave_externa": chaves.get(_norm(fundo)) or _nome_e_chave_fundo(fundo)[1],
    }
    (proventos if operacao == "PROVENTO" else movs).append(item)


def _simbolo_cripto(ativo: str, avisos: Optional[list]) -> str:
    """'XBT BITCOIN (XBT)' → 'XBT'. Sem parênteses, cai no slug do nome (e avisa)."""
    m = _RE_SIMBOLO_CRIPTO.search(ativo)
    if m:
        return m.group(1).upper()
    _avisar(avisos, "cripto_sem_simbolo", "CriptoAtivos",
            f"{ativo}: sem símbolo entre parênteses — chave caiu no nome inteiro")
    return _slug(ativo)


def _parse_cripto(
    wb, as_of: str, ignoradas: list[str], avisos: Optional[list] = None
) -> tuple[list[PosicaoParsed], list[dict], list[dict], list[dict]]:
    """
    (posições CRIPTO, proventos, movimentações, totais da aba). A aba pode ter só posição,
    só movimentações, ou ambas — posição ausente com movimentação presente é normal
    (posição zerada no fim do período). Atenção de layout: o rótulo é 'Valor Liquido'
    SEM acento nesta aba (irrelevante após _norm, registrado para o futuro leitor).
    """
    posicoes: list[PosicaoParsed] = []
    proventos: list[dict] = []
    movs: list[dict] = []
    totais: list[dict] = []
    ws = _aba(wb, "CriptoAtivos")
    if ws is None:
        return posicoes, proventos, movs, totais

    for tab in _tabelas(ws, "CriptoAtivos", ignoradas, avisos):
        t = _norm(tab.titulo)
        if t.startswith("posicao >"):
            i_ativo = _idx(tab, "ativo")
            i_qtde = _idx(tab, "quantidade", "qtde")
            i_preco = _idx(tab, "preco r$", "preco")
            i_bruto = _idx(tab, "valor bruto")
            i_liq = _idx(tab, "valor liquido")
            for row in tab.linhas:
                ativo = _txt(_cel(row, i_ativo))
                bruto = _num(_cel(row, i_bruto))
                liq = _num(_cel(row, i_liq))
                if not ativo or (bruto is None and liq is None):
                    ignoradas.append(_resumo_linha(f"CriptoAtivos/{tab.titulo}", row))
                    continue
                simbolo = _simbolo_cripto(ativo, avisos)
                valor = bruto if bruto is not None else liq
                if bruto is None:
                    _avisar(avisos, "cripto_sem_valor_bruto", "CriptoAtivos",
                            f"{simbolo}: sem Valor Bruto — usei o Valor Líquido")
                preco = _num(_cel(row, i_preco))
                qtde = _num(_cel(row, i_qtde))
                if qtde is None:
                    qtde, tipo = _derivar_quantidade(valor or 0.0, preco)
                    _avisar(avisos, tipo, "CriptoAtivos",
                            f"{simbolo}: quantidade ilegível — "
                            + ("derivada de valor/preço" if tipo == "quantidade_derivada" else "ficou 0.0"))
                posicoes.append(PosicaoParsed(
                    ticker=simbolo,
                    nome=ativo,
                    classe="CRIPTO",
                    quantidade=qtde,
                    preco_medio=None,
                    valor_mercado=valor or 0.0,
                    preco_fechamento=preco,
                    as_of=as_of,
                    chave_externa=f"CRIPTO:{simbolo}",
                ))
        elif t.startswith("movimentacao >"):
            i_data = _idx(tab, "data")
            i_ativo = _idx(tab, "ativo")
            i_desc = _idx(tab, "descricao")
            i_qtde = _idx(tab, "qtde", "quantidade")
            i_preco = _idx(tab, "preco r$", "preco")
            i_bruto = _idx(tab, "valor bruto")
            i_liq = _idx(tab, "valor liquido")
            for row in tab.linhas:
                descricao = _txt(_cel(row, i_desc))     # 'Venda' | 'Compra'
                ativo = _txt(_cel(row, i_ativo))
                if not descricao or not ativo:
                    if _tem_numero(row):
                        ignoradas.append(_resumo_linha(f"CriptoAtivos/{tab.titulo}", row))
                    continue
                simbolo = _simbolo_cripto(ativo, avisos)
                n = _norm(descricao)
                item = {
                    "data": _data(_cel(row, i_data)),
                    "transacao": descricao,
                    "ticker": None,     # mesma razão dos fundos: nada de chave B3 aqui
                    "quantidade": _num(_cel(row, i_qtde)),
                    "valor_bruto": _num(_cel(row, i_bruto)),
                    "valor_liquido": _num(_cel(row, i_liq)),
                    "classe": "CRIPTO",
                    "operacao": _operacao_rv(n),
                    "preco": _num(_cel(row, i_preco)),
                    "corretagem": None,
                    "amortizacao": False,
                    "nome": ativo,
                    "chave_externa": f"CRIPTO:{simbolo}",
                }
                (proventos if item["operacao"] == "PROVENTO" else movs).append(item)
        totais.extend(_totais_da_tabela(tab))

    return posicoes, proventos, movs, totais


# ---------------------------------------------------------------------------
# Conta Corrente (CAIXA) e Valores em Trânsito
# ---------------------------------------------------------------------------

def _parse_conta_corrente(wb, as_of: str) -> Optional[PosicaoParsed]:
    ws = _aba(wb, "Conta Corrente")
    if ws is None:
        return None

    for tab in _tabelas(ws):
        i_valor = _idx(tab, "valor financeiro")
        if i_valor is None:
            continue
        i_data = _idx(tab, "data")
        for row in reversed(tab.linhas):   # último saldo do período
            valor = _num(_cel(row, i_valor))
            if valor is None:
                continue
            return PosicaoParsed(
                ticker=None,
                nome="Conta corrente BTG",
                classe="CAIXA",
                quantidade=1.0,
                preco_medio=None,
                valor_mercado=valor,
                preco_fechamento=None,
                as_of=_data(_cel(row, i_data)) or as_of,
                chave_externa="CAIXA:BTG",
            )
    return None


def _parse_lancamentos_conta(
    wb, ignoradas: Optional[list] = None, avisos: Optional[list] = None
) -> tuple[Optional[float], list[dict], dict]:
    """
    Razão da conta corrente — a tabela 'Movimentações' da aba Conta Corrente.

    É a única fonte de aportes e resgates do mês: TED/PIX que entram ou saem da conta de
    investimento. Retorna (saldo_inicial, lançamentos, info), onde info leva os totais do
    próprio extrato e contadores para _checagem_conta.

    - 'Saldo Anterior' não é lançamento: vira saldo_inicial.
    - 'Saldo Final + Rendimento Provisionado...' é: traz o rendimento do saldo remunerado.
    - O sinal vem da diferença entre saldos consecutivos, quando ela confere com o valor
      da coluna. Nos extratos reais os débitos já vêm negativos na própria coluna; a
      regra fica como rede caso alguma release volte a imprimir só o valor absoluto.
    - Linha sem data herda a anterior; linha sem valor é descartada e só contada.
    """
    info: dict[str, Any] = {
        "encontrada": False,
        "saldo_final": None,
        "creditos_extrato": None,
        "debitos_extrato": None,
        "linhas_descartadas": 0,
        "sinal_divergente": 0,
    }
    ws = _aba(wb, "Conta Corrente")
    if ws is None:
        return None, [], info

    saldo_inicial: Optional[float] = None
    lancamentos: list[dict] = []
    # ignoradas/avisos entram só por aqui — _parse_conta_corrente varre a mesma aba e
    # registrar duas vezes duplicaria cada linha órfã.
    for tab in _tabelas(ws, "Conta Corrente", ignoradas, avisos):
        i_desc = _idx(tab, "descricao")
        i_mov = _idx(tab, "movimentacao")
        if i_desc is None or i_mov is None:
            continue
        info["encontrada"] = True
        i_data = _idx(tab, "data")
        i_saldo = _idx(tab, "saldo")

        ultima_data: Optional[str] = None
        saldo_corrente: Optional[float] = None   # saldo depois da linha anterior
        for row in tab.linhas:
            data = _data(_cel(row, i_data)) or ultima_data
            ultima_data = data
            texto = _cel(row, i_desc)
            valor = _num(_cel(row, i_mov))
            saldo = _num(_cel(row, i_saldo))

            if _norm(texto).startswith("saldo anterior"):
                if saldo is not None:
                    if saldo_inicial is None:
                        saldo_inicial = saldo
                    saldo_corrente = saldo
                continue
            if valor is None:
                info["linhas_descartadas"] += 1
                if saldo is not None:
                    saldo_corrente = saldo
                continue

            if saldo is not None and saldo_corrente is not None:
                diferenca = round(saldo - saldo_corrente, 2)
                if abs(abs(valor) - abs(diferenca)) <= 0.01:
                    valor = diferenca
                else:
                    info["sinal_divergente"] += 1

            lancamentos.append({
                "seq": len(lancamentos),
                "data": data,
                "descricao": sanitizar_descricao(texto, valor),
                "valor": round(valor, 2),
                "saldo": saldo,
            })
            if saldo is not None:
                saldo_corrente = saldo
            elif saldo_corrente is not None:
                saldo_corrente = round(saldo_corrente + valor, 2)

        for tot in tab.totais:
            preenchidas = _preenchidas(tot)
            rotulo = _norm(preenchidas[0]) if preenchidas else ""
            numero = _num(_cel(tot, i_mov))
            if numero is None:
                numero = next((n for n in (_num(c) for c in tot) if n is not None), None)
            # '-' no total = nada no mês
            if "credito" in rotulo:
                info["creditos_extrato"] = abs(numero) if numero is not None else 0.0
            elif "debito" in rotulo:
                info["debitos_extrato"] = abs(numero) if numero is not None else 0.0

        if saldo_corrente is not None:
            info["saldo_final"] = saldo_corrente

    return saldo_inicial, lancamentos, info


def _checagem_conta(
    saldo_inicial: Optional[float],
    lancamentos: list[dict],
    info: dict,
    caixa: Optional[PosicaoParsed],
    sumario: dict,
) -> dict:
    """
    O razão precisa fechar: saldo inicial + Σ lançamentos = saldo final, que é também a
    posição CAIXA; o saldo inicial é o Conta Corrente do Sumário anterior. Se não fecha,
    alguma linha não foi lida — e uma linha perdida pode ser justamente um aporte, então a
    rentabilidade do mês sai como provisória.
    """
    if not info.get("encontrada"):
        return {
            "ok": None,
            "encontrada": False,
            "aviso": (
                "A aba Conta Corrente não tem a tabela de movimentações — os aportes e "
                "resgates do mês não puderam ser lidos."
            ),
        }

    soma = round(sum(item["valor"] for item in lancamentos), 2)
    creditos = round(sum(item["valor"] for item in lancamentos if item["valor"] > 0), 2)
    debitos = round(sum(item["valor"] for item in lancamentos if item["valor"] < 0), 2)
    saldo_final = info.get("saldo_final")

    diferencas: dict[str, float] = {}
    if saldo_inicial is not None and saldo_final is not None:
        diferencas["diferenca"] = round(saldo_inicial + soma - saldo_final, 2)
    if info.get("creditos_extrato") is not None:
        diferencas["dif_creditos"] = round(creditos - info["creditos_extrato"], 2)
    if info.get("debitos_extrato") is not None:
        diferencas["dif_debitos"] = round(abs(debitos) - info["debitos_extrato"], 2)
    if caixa is not None and saldo_final is not None:
        diferencas["vs_caixa"] = round(saldo_final - caixa.valor_mercado, 2)
    cc_anterior = _sumario_mercado(sumario, "Conta Corrente", "anterior")
    if cc_anterior is not None and saldo_inicial is not None:
        diferencas["vs_sumario_anterior"] = round(saldo_inicial - cc_anterior, 2)

    ok = "diferenca" in diferencas and all(abs(v) <= TOLERANCIA_CONTA for v in diferencas.values())
    resultado: dict[str, Any] = {
        "ok": ok,
        "encontrada": True,
        "saldo_inicial": saldo_inicial,
        "soma_lancamentos": soma,
        "saldo_final": saldo_final,
        "creditos": creditos,
        "creditos_extrato": info.get("creditos_extrato"),
        "debitos": debitos,
        "debitos_extrato": info.get("debitos_extrato"),
        **diferencas,
        "linhas_descartadas": info.get("linhas_descartadas", 0),
        "sinal_divergente": info.get("sinal_divergente", 0),
        "tolerancia": TOLERANCIA_CONTA,
    }
    if not ok:
        resultado["aviso"] = (
            "O razão da conta corrente não fecha com os saldos do extrato — algum "
            "lançamento pode não ter sido lido. A rentabilidade do mês fica provisória."
        )
    return resultado


def _parse_valores_transito(
    wb, ignoradas: Optional[list] = None, avisos: Optional[list] = None
) -> tuple[list[dict], list[dict]]:
    """(itens em trânsito, totais da aba)."""
    ws = _aba(wb, "Valores em Trânsito")
    if ws is None:
        return [], []

    out: list[dict] = []
    totais: list[dict] = []
    for tab in _tabelas(ws, "Valores em Transito", ignoradas, avisos):
        i_data = _idx(tab, "data liquidacao", "data")
        i_desc = _idx(tab, "descricao")
        i_valor = _idx(tab, "valor r$", "valor")
        if i_valor is None:
            continue
        for row in tab.linhas:
            valor = _num(_cel(row, i_valor))
            if valor is None:
                continue
            out.append({
                "data_liquidacao": _data(_cel(row, i_data)),
                "descricao": sanitizar_descricao(_cel(row, i_desc), valor),
                "valor": valor,
                "origem": tab.titulo,
            })
        totais.extend(_totais_da_tabela(tab))
    return out, totais


# ---------------------------------------------------------------------------
# Checksum contra a aba Sumario
# ---------------------------------------------------------------------------

_CLASSES_RV = {"ACAO", "ETF", "FII", "BDR"}
_CLASSES_RF = {"TESOURO", "RF"}


def _checagem(
    posicoes: list[PosicaoParsed], sumario: dict, transito: float, aluguel: Sequence[dict] = ()
) -> dict:
    def _soma(classes: set[str]) -> float:
        return round(sum(p.valor_mercado for p in posicoes if p.classe in classes), 2)

    # O Sumário do BTG soma o resultado acumulado do aluguel de ações ao total de Renda
    # Variável (e portanto ao Total) — mas aluguel não é posição (ver _parse_aluguel).
    # Sem este desconto a checagem de RV acusava divergência permanente em quem aluga.
    rendimento_aluguel = round(
        sum(a.get("resultado_acumulado_liquido") or 0.0 for a in aluguel), 2
    )

    parseado = {
        "Renda Variável": _soma(_CLASSES_RV),
        "Renda Fixa": _soma(_CLASSES_RF),
        "Fundos de Investimento": _soma({"FUNDO"}),
        "CriptoAtivos": _soma({"CRIPTO"}),
        "Conta Corrente": _soma({"CAIXA"}),
    }

    comparacoes: list[dict] = []
    for mercado, valor in parseado.items():
        esperado = _sumario_mercado(sumario, mercado)
        if esperado is None:
            # Mercado ausente do Sumário ou com '-' (classe zerada no fim do período):
            # nada a comparar aqui — a cobertura é assunto do validador (V8).
            continue
        if mercado == "Renda Variável":
            esperado -= rendimento_aluguel
        comparacoes.append({
            "mercado": mercado,
            "parseado": valor,
            "sumario": round(esperado, 2),
            "diferenca": round(valor - esperado, 2),
        })

    total_parseado = round(sum(p.valor_mercado for p in posicoes), 2)
    total_sumario = ((sumario.get("atual") or {}).get("total") or {}).get("bruto")
    # O total do Sumário inclui Valores em Trânsito (que não são posição) e o resultado
    # do aluguel (idem).
    esperado_total = (
        round(total_sumario - transito - rendimento_aluguel, 2)
        if total_sumario is not None else None
    )
    diferenca_total = round(total_parseado - esperado_total, 2) if esperado_total is not None else None

    ok = (
        diferenca_total is not None
        and abs(diferenca_total) <= TOLERANCIA_CHECKSUM
        and all(abs(c["diferenca"]) <= TOLERANCIA_CHECKSUM for c in comparacoes)
    )

    resultado = {
        "ok": ok,
        "total_parseado": total_parseado,
        "total_esperado": esperado_total,
        "diferenca": diferenca_total,
        "por_mercado": comparacoes,
        "valores_em_transito_excluidos": round(transito, 2),
        "rendimento_aluguel_excluido": rendimento_aluguel,
        "tolerancia": TOLERANCIA_CHECKSUM,
    }
    if esperado_total is None:
        resultado["aviso"] = "Aba 'Sumario' sem total — não foi possível conferir o import."
    elif not ok:
        resultado["aviso"] = (
            f"O total das posições lidas (R$ {total_parseado:,.2f}) não bate com o "
            f"Sumário do extrato (R$ {esperado_total:,.2f}). Confira antes de gravar."
        )
    return resultado


# ---------------------------------------------------------------------------
# Parser principal
# ---------------------------------------------------------------------------

def parse_btg_xlsx(conteudo: bytes) -> ExtratoParsed:
    """
    Parseia o extrato XLSX do BTG. Levanta ExtratoParseError se o arquivo não for
    legível ou não tiver data de referência.
    """
    try:
        wb = load_workbook(BytesIO(conteudo), data_only=True)
    except Exception as e:
        raise ExtratoParseError(f"Não consegui abrir o arquivo como XLSX: {e}") from e

    try:
        data_ref, meta = _parse_capa(wb)
        ignoradas: list[str] = []
        avisos: list[dict] = []

        sumario = _parse_sumario(wb)
        pos_rv, proventos, movimentacoes, aluguel, blocos_rv, totais_rv = _parse_renda_variavel(
            wb, data_ref, ignoradas, avisos
        )
        pos_rf, lotes_rf, totais_rf = _parse_renda_fixa(wb, data_ref, ignoradas, avisos)
        pos_fundos, prov_fundos, movs_fundos, totais_fundos = _parse_fundos(
            wb, data_ref, ignoradas, avisos
        )
        pos_cripto, prov_cripto, movs_cripto, totais_cripto = _parse_cripto(
            wb, data_ref, ignoradas, avisos
        )
        caixa = _parse_conta_corrente(wb, data_ref)
        saldo_inicial_conta, lancamentos, info_conta = _parse_lancamentos_conta(
            wb, ignoradas, avisos
        )
        transito, totais_transito = _parse_valores_transito(wb, ignoradas, avisos)
    finally:
        wb.close()

    posicoes = pos_rv + pos_rf + pos_fundos + pos_cripto + ([caixa] if caixa else [])
    if not posicoes:
        raise ExtratoParseError(
            "Nenhuma posição encontrada no arquivo. Confira se o XLSX é o extrato da "
            "conta de investimento do BTG (abas 'Renda Variavel' / 'Renda Fixa')."
        )
    proventos = proventos + prov_fundos + prov_cripto
    movimentacoes = movimentacoes + movs_fundos + movs_cripto

    total_transito = sum(v["valor"] for v in transito)
    # `ok` continua sendo a conferência das POSIÇÕES contra o Sumário; as conferências do
    # razão e das movimentações vão em chaves próprias, sem mudar o que o import bloqueia.
    checagem = _checagem(posicoes, sumario, total_transito, aluguel)
    checagem["conta_corrente"] = _checagem_conta(
        saldo_inicial_conta, lancamentos, info_conta, caixa, sumario
    )
    checagem["movimentacao_rv"] = {
        "ok": all(b["ok"] for b in blocos_rv),
        "blocos": blocos_rv,
    }

    extrato = ExtratoParsed(
        posicoes=posicoes,
        data_referencia=data_ref,
        sumario=sumario,
        proventos=proventos,
        movimentacoes=movimentacoes,
        aluguel=aluguel,
        valores_em_transito=transito,
        linhas_ignoradas=ignoradas,
        checagem=checagem,
        saldo_inicial_conta=saldo_inicial_conta,
        lancamentos_conta=lancamentos,
        lotes_rf=lotes_rf,
        totais_abas={
            "renda_variavel": totais_rv,
            "renda_fixa": totais_rf,
            "fundos": totais_fundos,
            "cripto": totais_cripto,
            "valores_em_transito": totais_transito,
        },
        avisos_parser=avisos,
    )
    extrato.sumario = {**sumario, "meta": meta} if sumario else {"meta": meta}

    logger.info(
        "parse_btg_xlsx: %d posicoes, ref=%s, checksum_ok=%s, %d linha(s) ignorada(s), "
        "%d aviso(s) do parser, %d lancamento(s) na conta (razao_ok=%s), %d lote(s) de RF",
        len(posicoes), data_ref, checagem.get("ok"), len(ignoradas),
        len(avisos), len(lancamentos), checagem["conta_corrente"].get("ok"), len(lotes_rf),
    )
    return extrato
