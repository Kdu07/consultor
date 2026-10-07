"""
Gerador de extratos XLSX sintéticos no formato do BTG Pactual.

Substitui a antiga fixture de arquivo (cópia anonimizada de um extrato real): aqui o
workbook inteiro nasce do zero, com a MESMA FORMA de um extrato de verdade — mesmas abas,
mesmos títulos de bloco, mesmos cabeçalhos, mesmas contagens de linhas — e VALORES
inventados. Nenhum número vem de extrato real; os tickers são papéis públicos da B3 e os
textos imitam o vocabulário do BTG.

O gerador produz os invariantes do extrato real COERENTES POR CONSTRUÇÃO:

  - Σ linhas de mercado do Sumário = Total, nas quatro colunas (bruto/líquido × ini/fim);
  - Sumário de Renda Variável (fim, bruto) = Σ 'Total em ...' dos blocos de posição
    + resultado acumulado do aluguel de ações (o aluguel NÃO é posição);
  - Sumário de Renda Fixa = Σ 'Total' das seções de Posição, bruto e líquido;
  - o razão da conta corrente fecha linha a linha: 'Saldo Anterior' = Conta Corrente do
    Sumário anterior, último saldo = posição da conta = Sumário atual, Σ movimentos =
    Total de Créditos + Total de Débitos;
  - Σ Valores em Trânsito = Total da aba = Sumário;
  - qtde × preço = saldo em toda linha de posição (RV exato; RF via lotes).

Os knobs da EspecExtrato injetam exatamente um defeito de cada vez (ver a dataclass), para
os testes do validador programarem contra severidades conhecidas.

As constantes *_FIXTURE no fim do módulo são derivadas das mesmas tabelas que o builder
usa — asserts de teste devem referenciá-las em vez de repetir números, para que um
reequilíbrio futuro dos valores sintéticos não exija nova migração.

O nome não começa com test_: o pytest não coleta este módulo.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

# ---------------------------------------------------------------------------
# Especificação
# ---------------------------------------------------------------------------

_NULO = "-"   # como o BTG escreve célula sem valor


@dataclass
class EspecExtrato:
    """
    Um extrato sintético. Os defaults reproduzem a fixture padrão (mês civil completo,
    carteira coerente). Cada knob injeta um desvio visto em extratos reais:

    decorar_tickers      extrato de período aberto: 'BBAS3*' na RV, 'LFT* ' na RF,
                         'Renda Variável*' / 'Renda Fixa*' no Sumário e a nota de rodapé
                         '* Seção com ativos calculados...' nas abas.
    descontinuidade_cc   salto no saldo do razão sem lançamento correspondente (o BTG real
                         já exibiu R$ 0,14): o razão deixa de fechar por exatamente esse
                         valor; a posição da conta e o Sumário seguem o saldo final saltado.
    saldo_anterior_delta 'Saldo Anterior' do razão ≠ Conta Corrente do Sumário anterior
                         por este delta (BTG real: ±R$ 0,02–0,05). O razão continua
                         fechando internamente; só a amarração com o Sumário anterior
                         diverge, e o saldo final carrega o delta.
    desalinhar_sumario   {mercado: delta} somado ao saldo bruto FIM daquele mercado E ao
                         Total (o Sumário continua internamente consistente; a divergência
                         fica entre Sumário e os blocos da aba — alvo dos checks V2/V3/V5/V6).
    desalinhar_total     delta somado SÓ à linha Total (fim, bruto): quebra Σ mercados =
                         Total — alvo do check V1.
    sumario_extra        {mercado: valor_fim} de mercados SEM aba correspondente
                         (ini = fim = valor): alvo do check V8 de cobertura.
    com_fundos           aba 'Fundos' (Posição > Portfólio de fundos, registro em DUAS
                         linhas — nome com 'Classe CNPJ:' numa linha, dados na seguinte —
                         e Movimentação > {fundo}) + linha 'Fundos de Investimento' no Sumário.
    com_cripto           aba 'CriptoAtivos' + linha 'CriptoAtivos' no Sumário. Atenção à
                         grafia real: 'Valor Liquido R$' SEM acento nesta aba.
    cripto_so_movimentacao  aba CriptoAtivos sem o bloco de posição (só movimentação) e
                         Sumário com '-' nas colunas fim — posição zerada no período.
    lancamentos_extras   [(date, descrição, valor com sinal)] acrescentados ao razão em
                         ordem de data; o saldo, os totais, a posição da conta e o Sumário
                         acompanham. É o lugar de PIX/TED com nome/CPF FICTÍCIOS para os
                         testes de privacidade.
    aluguel_resultado    resultado acumulado líquido do aluguel de ações. > R$ 1,00 de
                         propósito: com o parser que desconta o aluguel do Sumário a
                         checagem fecha; um parser que não desconta estoura a tolerância.
    transito_ini         Valores em Trânsito do Sumário anterior (default: igual ao fim).
    """
    inicio: date = date(2026, 7, 1)
    fim: date = date(2026, 7, 31)
    decorar_tickers: bool = False
    descontinuidade_cc: float = 0.0
    saldo_anterior_delta: float = 0.0
    desalinhar_sumario: dict[str, float] = field(default_factory=dict)
    desalinhar_total: float = 0.0
    sumario_extra: dict[str, float] = field(default_factory=dict)
    com_fundos: bool = False
    com_cripto: bool = False
    cripto_so_movimentacao: bool = False
    lancamentos_extras: list[tuple[date, str, float]] = field(default_factory=list)
    aluguel_resultado: float = 1.37
    transito_ini: float | None = None


# ---------------------------------------------------------------------------
# Carteira sintética padrão (valores INVENTADOS; tickers públicos da B3)
# ---------------------------------------------------------------------------

# (ticker, nome como o BTG escreve — espaços múltiplos de propósito, qtde, preço, preço médio)
_ACOES = (
    ("BBAS3", "BRASIL      ON      NM", 100, 20.00, 24.00),
    ("ISAE4", "ISA ENERGIA PN      N1", 100, 25.00, 20.00),
    ("ITUB4", "ITAUUNIBANCOPN      N1", 50, 40.00, 38.00),
    ("TAEE11", "TAESA       UNT     N2", 100, 36.00, 30.00),
    ("VALE3", "VALE        ON      NM", 10, 70.00, 65.00),
)
_ETFS = (("IVVB11", "ISHARE SP500CI", 5, 400.00, 390.00),)
# (ticker, nome, tipo, qtde, preço, preço médio)
_FIIS = (
    ("HGCR11", "FII HGCR PAXCI", "FII", 40, 100.00, 98.00),
    ("KNCR11", "FII KINEA RICI", "FII", 60, 110.00, 105.00),
    ("RBRR11", "FII RBRR PAXCI", "FII", 20, 80.00, 90.00),
)

_EMISSOR_TD = "BACEN-BANCO CENTRAL DO BRASIL - RJ"


@dataclass(frozen=True)
class _LoteTD:
    sigla: str
    emissao: date
    vencimento: date
    aquisicao: date
    taxa: str
    qtde: float
    preco_compra: float
    preco_atual: float
    ir: float

    @property
    def valor_compra(self) -> float:
        return round(self.qtde * self.preco_compra, 2)

    @property
    def saldo(self) -> float:
        return round(self.qtde * self.preco_atual, 2)

    @property
    def liquido(self) -> float:
        return round(self.saldo - self.ir, 2)


# 5 posições de Tesouro (LFT 2031/2028, LTN 2028/2029, NTNB-P 2029) em 8 lotes — a NTNB-P
# tem 4 lotes, como num extrato real de quem aporta aos poucos no mesmo título.
_LOTES_TD = (
    _LoteTD("LFT", date(2025, 1, 8), date(2031, 3, 1), date(2026, 2, 19), "SELIC + 0,10%", 0.10, 18000.00, 19000.00, 10.00),
    _LoteTD("LFT", date(2022, 1, 5), date(2028, 3, 1), date(2026, 1, 15), "SELIC + 0,04%", 0.20, 17500.00, 19200.00, 40.00),
    _LoteTD("LTN", date(2024, 1, 5), date(2028, 1, 1), date(2025, 11, 10), "13,00% a.a.", 2.00, 750.00, 800.00, 15.00),
    _LoteTD("LTN", date(2022, 2, 11), date(2029, 1, 1), date(2026, 3, 25), "13,50% a.a.", 5.00, 700.00, 740.00, 25.00),
    _LoteTD("NTNB-P", date(2022, 12, 19), date(2029, 5, 15), date(2024, 12, 2), "IPCA + 7,16%", 0.30, 3100.00, 3600.00, 25.00),
    _LoteTD("NTNB-P", date(2022, 12, 19), date(2029, 5, 15), date(2025, 4, 24), "IPCA + 7,58%", 0.50, 3200.00, 3600.00, 30.00),
    _LoteTD("NTNB-P", date(2022, 12, 19), date(2029, 5, 15), date(2025, 8, 11), "IPCA + 7,69%", 0.70, 3300.00, 3600.00, 45.00),
    _LoteTD("NTNB-P", date(2022, 12, 19), date(2029, 5, 15), date(2025, 11, 10), "IPCA + 7,89%", 0.50, 3400.00, 3600.00, 20.00),
)
_ORDEM_TD = ("LFT", "LTN", "NTNB-P")
_TAXA_MEDIA_TD = {
    ("LFT", "2031-03-01"): "SELIC + 0,10%",
    ("LFT", "2028-03-01"): "SELIC + 0,04%",
    ("LTN", "2028-01-01"): "13,00% a.a.",
    ("LTN", "2029-01-01"): "13,50% a.a.",
    ("NTNB-P", "2029-05-15"): "IPCA + 7,62%",
}

# Proventos do mês: (dias após o início, ticker, transação, bloco de RV, qtde, bruto, líquido).
# O valor do razão TEM de ser igual ao valor líquido da movimentação de RV — é assim que o
# motor de desempenho reconhece que o crédito do razão e o provento da RV são o mesmo dinheiro.
_PROVENTOS = (
    (0, "ITUB4", "JUROS S/CAPITAL", "acoes", 50, 1.00, 0.85),
    (12, "KNCR11", "RENDIMENTO", "fiis", 60, 66.00, 66.00),
    (13, "HGCR11", "RENDIMENTO", "fiis", 40, 44.00, 44.00),
    (15, "RBRR11", "RENDIMENTO", "fiis", 20, 21.00, 21.00),
)

# Aluguel de ações: TAEE11 doada — informativo, NUNCA posição (o papel já está em Ações).
_ALUGUEL_TICKER = "TAEE11"
_ALUGUEL_PRECO_REF = 36.10
_ALUGUEL_QTDE = 100
_ALUGUEL_TAXA_ANO = 0.15
_ALUGUEL_IR = 0.02

# Rendimento do saldo remunerado — a última linha do razão.
_RENDIMENTO_SALDO = 0.07
_DESC_SALDO_FINAL = "Saldo Final + Rendimento Provisionado de Saldo Remunerado"

# Valores em trânsito: (dias após o fim, descrição, valor)
_TRANSITO = (
    (3, "JUROS S/ CAPITAL - ITAUUNIBANCOPN      N1", 1.00),
    (26, "JUROS S/ CAPITAL - TAESA       UNT     N2", 45.00),
    (31, "JUROS S/ CAPITAL - ITAUUNIBANCOPN      N1", 19.00),
)

# Sumário do período anterior (inventado; só precisa somar no Total)
_RV_INI = 24800.00
_RF_INI_BRUTO = 18100.00
_RF_INI_LIQUIDO = 17900.00
_CAIXA_INI = 150.00

# Aba Fundos (opcional) — fundo e CNPJ fictícios e óbvios
_FUNDO_NOME = "FUNDO EXEMPLO RENDA FIXA CP RL"
_FUNDO_CNPJ = "00.000.000/0001-00"
_FUNDO_LIQ_ANTERIOR = 3000.00
_FUNDO_QTDE_COTAS = 25.00
_FUNDO_COTACAO = 124.00            # 25 × 124,00 = 3.100,00 bruto
_FUNDO_IR = 20.00                  # líquido 3.080,00
_FUNDO_INI = (3000.00, 2980.00)    # (bruto, líquido) no Sumário anterior
_FUNDO_APLICACAO = (4.00, 123.00)  # (cotas, valor da cota) → aplicação de 492,00 no mês

# Aba CriptoAtivos (opcional)
_CRIPTO_ATIVO = "XBT BITCOIN (XBT)"
_CRIPTO_QTDE = 0.0015
_CRIPTO_PRECO = 400000.00          # 0,0015 × 400.000,00 = 600,00
_CRIPTO_INI = 520.00
_CRIPTO_MOV = (2, "XBT", "COMPRA DE CRIPTOATIVOS", 0.0005, 400000.00, 200.00)


# ---------------------------------------------------------------------------
# Derivações (as mesmas somas que o extrato real exibe prontas)
# ---------------------------------------------------------------------------

def _r2(v: float) -> float:
    return round(v, 2)


def _soma_acoes() -> float:
    return _r2(sum(q * p for _, _, q, p, _ in _ACOES))


def _soma_etfs() -> float:
    return _r2(sum(q * p for _, _, q, p, _ in _ETFS))


def _soma_fiis() -> float:
    return _r2(sum(q * p for _, _, _, q, p, _ in _FIIS))


def _posicoes_td() -> list[dict]:
    """Agrega os lotes por (sigla, vencimento) na ordem das seções do extrato."""
    out: list[dict] = []
    for sigla in _ORDEM_TD:
        por_venc: dict[date, dict] = {}
        for lote in _LOTES_TD:
            if lote.sigla != sigla:
                continue
            acc = por_venc.setdefault(lote.vencimento, {
                "sigla": sigla, "emissao": lote.emissao, "vencimento": lote.vencimento,
                "preco": lote.preco_atual, "qtde": 0.0, "saldo": 0.0, "ir": 0.0,
            })
            acc["qtde"] = _r2(acc["qtde"] + lote.qtde)
            acc["saldo"] = _r2(acc["saldo"] + lote.saldo)
            acc["ir"] = _r2(acc["ir"] + lote.ir)
        for pos in por_venc.values():
            pos["liquido"] = _r2(pos["saldo"] - pos["ir"])
            pos["taxa"] = _TAXA_MEDIA_TD[(sigla, pos["vencimento"].isoformat())]
            out.append(pos)
    return out


def _totais_secao_td(sigla: str) -> tuple[float, float, float, float]:
    """(valor compra, saldo bruto, IR, saldo líquido) da seção da sigla."""
    lotes = [l for l in _LOTES_TD if l.sigla == sigla]
    return (
        _r2(sum(l.valor_compra for l in lotes)),
        _r2(sum(l.saldo for l in lotes)),
        _r2(sum(l.ir for l in lotes)),
        _r2(sum(l.liquido for l in lotes)),
    )


def _rf_totais() -> tuple[float, float]:
    bruto = _r2(sum(l.saldo for l in _LOTES_TD))
    liquido = _r2(sum(l.liquido for l in _LOTES_TD))
    return bruto, liquido


def _transito_total() -> float:
    return _r2(sum(v for _, _, v in _TRANSITO))


def _dt(d: date) -> datetime:
    return datetime(d.year, d.month, d.day)


def _dia(spec: EspecExtrato, offset: int) -> date:
    """Data dentro do período: início + offset, nunca depois do fim."""
    return min(spec.inicio + timedelta(days=offset), spec.fim)


def _razao(spec: EspecExtrato) -> tuple[float, list[tuple[date, str, float, float]], float]:
    """
    (saldo anterior, linhas [(data, descrição, valor, saldo)], saldo final).
    Proventos + lançamentos extras em ordem de data; 'Saldo Final + Rendimento...' por
    último. A descontinuidade (se houver) entra no meio: o saldo salta sem lançamento.
    """
    nomes = {t: n for t, n, *_ in _ACOES}
    nomes.update({t: n for t, n, *_ in _ETFS})
    nomes.update({t: n for t, n, _, _, _, _ in _FIIS})

    eventos: list[tuple[date, str, float]] = []
    for offset, ticker, transacao, _, _, _, liquido in _PROVENTOS:
        if transacao.startswith("JUROS"):
            desc = f"JUROS S/ CAPITAL - À VISTA s/ {nomes[ticker]} - {ticker}"
        else:
            desc = f"RENDIMENTOS - À VISTA s/ {nomes[ticker]} - {ticker}"
        eventos.append((_dia(spec, offset), desc, liquido))
    eventos.extend(
        (d.date() if isinstance(d, datetime) else d, desc, _r2(v))
        for d, desc, v in spec.lancamentos_extras
    )
    eventos.sort(key=lambda e: e[0])
    eventos.append((spec.fim, _DESC_SALDO_FINAL, _RENDIMENTO_SALDO))

    salto_em = len(eventos) // 2
    saldo = _r2(_CAIXA_INI + spec.saldo_anterior_delta)
    saldo_anterior = saldo
    linhas: list[tuple[date, str, float, float]] = []
    for i, (d, desc, valor) in enumerate(eventos):
        if spec.descontinuidade_cc and i == salto_em:
            saldo = _r2(saldo + spec.descontinuidade_cc)   # salto sem lançamento
        saldo = _r2(saldo + valor)
        linhas.append((d, desc, valor, saldo))
    return saldo_anterior, linhas, saldo


def _linhas_sumario(spec: EspecExtrato) -> list[tuple[str, object, object, object, object]]:
    """[(mercado, bruto_ini, liq_ini, bruto_fim, liq_fim)] + linha Total, já desalinhado."""
    rv_fim = _r2(_soma_acoes() + _soma_etfs() + _soma_fiis() + spec.aluguel_resultado)
    rv_fim_liq = _r2(rv_fim - _ALUGUEL_IR)
    rf_bruto, rf_liquido = _rf_totais()
    _, _, caixa_fim = _razao(spec)
    transito = _transito_total()
    transito_ini = transito if spec.transito_ini is None else _r2(spec.transito_ini)

    decor = "*" if spec.decorar_tickers else ""
    linhas: list[tuple[str, object, object, object, object]] = [
        (f"Renda Variável{decor}", _RV_INI, _RV_INI, rv_fim, rv_fim_liq),
        (f"Renda Fixa{decor}", _RF_INI_BRUTO, _RF_INI_LIQUIDO, rf_bruto, rf_liquido),
    ]
    if spec.com_cripto:
        cripto_bruto = _r2(_CRIPTO_QTDE * _CRIPTO_PRECO)
        if spec.cripto_so_movimentacao:
            linhas.append(("CriptoAtivos", _CRIPTO_INI, _CRIPTO_INI, _NULO, _NULO))
        else:
            linhas.append(("CriptoAtivos", _CRIPTO_INI, _CRIPTO_INI, cripto_bruto, cripto_bruto))
    linhas.append(("Conta Corrente", _CAIXA_INI, _CAIXA_INI, caixa_fim, caixa_fim))
    linhas.append(("Valores em Trânsito", transito_ini, transito_ini, transito, transito))
    if spec.com_fundos:
        fundo_bruto = _r2(_FUNDO_QTDE_COTAS * _FUNDO_COTACAO)
        linhas.append(("Fundos de Investimento", _FUNDO_INI[0], _FUNDO_INI[1],
                       fundo_bruto, _r2(fundo_bruto - _FUNDO_IR)))
    for mercado, valor in spec.sumario_extra.items():
        linhas.append((mercado, _r2(valor), _r2(valor), _r2(valor), _r2(valor)))

    ajustadas = []
    for nome, b_ini, l_ini, b_fim, l_fim in linhas:
        delta = spec.desalinhar_sumario.get(nome.rstrip("*"))
        if delta and isinstance(b_fim, float):
            b_fim = _r2(b_fim + delta)
        ajustadas.append((nome, b_ini, l_ini, b_fim, l_fim))

    def _soma(i: int) -> float:
        return _r2(sum(l[i] for l in ajustadas if isinstance(l[i], (int, float))))

    total = ("Total", _soma(1), _soma(2), _r2(_soma(3) + spec.desalinhar_total), _soma(4))
    return ajustadas + [total]


# ---------------------------------------------------------------------------
# Escrita do workbook
# ---------------------------------------------------------------------------

class _Folha:
    """Cursor de escrita: coluna A sempre vazia, como no extrato real."""

    def __init__(self, wb: Workbook, titulo: str):
        self.ws: Worksheet = wb.create_sheet(titulo)
        self.r = 1

    def pular(self, n: int = 1) -> None:
        self.r += n

    def linha(self, *valores: object, col: int = 2) -> None:
        for i, v in enumerate(valores):
            if v is None:
                continue
            self.ws.cell(self.r, col + i).value = _dt(v) if isinstance(v, date) and not isinstance(v, datetime) else v
        self.r += 1


def _rodape_decoracao(f: _Folha, spec: EspecExtrato) -> None:
    if spec.decorar_tickers:
        f.pular()
        f.linha("* Seção com ativos calculados em data anterior à data da geração do extrato")


def _aba_capa(wb: Workbook, spec: EspecExtrato) -> None:
    f = _Folha(wb, "Capa")
    f.pular(3)
    f.linha(None, "Extrato da Conta Investimento")
    f.pular(7)
    f.linha(None, f"Período de {spec.inicio:%d/%m/%y} a {spec.fim:%d/%m/%y}")
    f.pular()
    f.linha(None, f"Emitido em {spec.fim + timedelta(days=4):%d/%m/%y} 23:29")
    f.pular(2)
    f.linha(None, "FULANO DE TAL DA SILVA")          # titular fictício
    f.pular(3)
    f.linha(None, "Banco: BTG Pactual")
    f.linha(None, "Agência: 9999")
    f.linha(None, "Conta Controle: 000000000")
    f.linha(None, "CPF: 000.000.000-00")


def _aba_sumario(wb: Workbook, spec: EspecExtrato) -> None:
    f = _Folha(wb, "Sumario")
    f.pular()
    f.linha("Sumário")
    f.pular()
    f.linha("Distribuição")
    f.pular()
    anterior = spec.inicio - timedelta(days=1)
    f.linha(
        "Mercados",
        f"Saldo Bruto R$ {anterior:%d/%m/%y}", f"Saldo Líquido R$ {anterior:%d/%m/%y}",
        f"Saldo Bruto R$ {spec.fim:%d/%m/%y}", f"Saldo Líquido R$ {spec.fim:%d/%m/%y}",
    )
    for nome, b_ini, l_ini, b_fim, l_fim in _linhas_sumario(spec):
        f.linha(nome, b_ini, l_ini, b_fim, l_fim)


def _aba_renda_fixa(wb: Workbook, spec: EspecExtrato) -> None:
    f = _Folha(wb, "Renda Fixa")
    sigla_de = (lambda s: f"{s}* ") if spec.decorar_tickers else (lambda s: s)
    f.pular()
    f.linha("Renda Fixa")
    f.pular()
    f.linha("Posições")
    f.pular()

    cabecalho_pos = (
        "Emissor", "Ativo", "Emissão", "Vencimento", "Liquidez",
        "Dias de carência para liquidez", "Data inicial de liquidez",
        "Taxa Média Ponderada", "Quantidade", "Preço R$", "Saldo Bruto R$",
        "IR R$", "IOF R$", "Saldo Líquido R$",
    )
    posicoes = _posicoes_td()
    for sigla in _ORDEM_TD:
        f.linha(f"Posição > TESOURO DIRETO - {sigla}")
        f.linha(*cabecalho_pos)
        secao = [p for p in posicoes if p["sigla"] == sigla]
        for p in secao:
            f.linha(_EMISSOR_TD, sigla_de(sigla), p["emissao"], p["vencimento"], "Não",
                    _NULO, _NULO, p["taxa"], p["qtde"], p["preco"], p["saldo"],
                    p["ir"], _NULO, p["liquido"])
        f.linha("Total", None, None, None, None, None, None, None, None, None,
                _r2(sum(p["saldo"] for p in secao)), _r2(sum(p["ir"] for p in secao)),
                _NULO, _r2(sum(p["liquido"] for p in secao)))
        f.pular()

    f.pular()
    f.linha("Posições Detalhadas")
    f.pular()
    cabecalho_det = (
        "Ativo", "Emissão", "Vencimento", "Aquisição", "Liquidez",
        "Dias de carência para liquidez", "Data inicial de liquidez", "Taxa Compra",
        "Quantidade", "Preço Compra R$", "Valor Compra R$", "Preço R$", "Saldo Bruto R$",
        "IR R$", "IOF R$", "Saldo Líquido R$",
    )
    for sigla in _ORDEM_TD:
        f.linha(f"Detalhamento > TESOURO DIRETO - {sigla} | {_EMISSOR_TD}")
        f.linha(*cabecalho_det)
        for lote in (l for l in _LOTES_TD if l.sigla == sigla):
            f.linha(sigla_de(sigla), lote.emissao, lote.vencimento, lote.aquisicao, "Não",
                    _NULO, _NULO, lote.taxa, lote.qtde, lote.preco_compra,
                    lote.valor_compra, lote.preco_atual, lote.saldo, lote.ir, _NULO,
                    lote.liquido)
        compra, saldo, ir, liquido = _totais_secao_td(sigla)
        f.linha("Total", None, None, None, None, None, None, None, None, None,
                compra, None, saldo, ir, _NULO, liquido)
        f.pular()

    f.pular()
    f.linha("Posição Consolidada Por Emissor")
    f.pular()
    f.linha("Posição Consolidada Por Emissor > ")
    f.linha("Emissor", "Saldo Bruto R$")
    f.linha(_EMISSOR_TD, _rf_totais()[0])
    _rodape_decoracao(f, spec)


def _aba_renda_variavel(wb: Workbook, spec: EspecExtrato) -> None:
    f = _Folha(wb, "Renda Variavel")
    tk = (lambda t: f"{t}*") if spec.decorar_tickers else (lambda t: t)
    f.pular()
    f.linha("Renda Variável")
    f.pular()

    # Posição > Ações
    f.linha("Posição")
    f.linha("Posição > Ações")
    f.linha("Código", "Ação", "Qtde.", "Preço Fechamento R$", "Preço Médio R$", "Saldo Bruto R$")
    for ticker, nome, qtde, preco, medio in _ACOES:
        f.linha(tk(ticker), nome, qtde, preco, medio, _r2(qtde * preco))
    f.linha("Total em Ações R$", None, None, None, None, _soma_acoes())
    f.pular(2)

    # Movimentação > Ações (proventos do bloco)
    f.linha("Movimentação")
    f.linha("Movimentação > Ações")
    f.linha("Data", "Transação", "Código", "Qtde.", "Preço R$", "Valor Bruto R$",
            "Corretagem e Emolumentos R$", "Valor Líquido R$")
    proventos_acoes = [p for p in _PROVENTOS if p[3] == "acoes"]
    for offset, ticker, transacao, _, qtde, bruto, liquido in proventos_acoes:
        f.linha(_dia(spec, offset), transacao, tk(ticker), qtde, _NULO, bruto, _NULO, liquido)
    f.linha("Total de Compras", None, None, None, None, _NULO, _NULO, _NULO)
    f.linha("Total de Vendas", None, None, None, None, _NULO, _NULO, _NULO)
    f.linha("Total de Proventos", None, None, None, None,
            _r2(sum(p[5] for p in proventos_acoes)), None, _r2(sum(p[6] for p in proventos_acoes)))
    f.pular(2)

    # Posição > Ações | Aluguel — papel doado; já contado na posição de Ações
    contratado = _r2(_ALUGUEL_QTDE * _ALUGUEL_PRECO_REF)
    f.linha("Posição")
    f.linha("Posição > Ações | Aluguel")
    f.linha("Código", "Qtde.", "Posição", "Preço de Referência R$", "Valor Contratado R$",
            "Data Operação", "Data Vencimento", "Taxa Ano %", "IR R$", "Valor Repasse R$",
            "Result. Acum. Líq.  R$")
    f.linha(tk(_ALUGUEL_TICKER), _ALUGUEL_QTDE, "Doador", _ALUGUEL_PRECO_REF, contratado,
            spec.fim - timedelta(days=4), spec.fim + timedelta(days=26), _ALUGUEL_TAXA_ANO,
            _ALUGUEL_IR, _NULO, _r2(spec.aluguel_resultado))
    f.linha("Total em Aluguel de Ações R$", None, None, None, contratado, None, None, None,
            None, None, _r2(spec.aluguel_resultado))
    f.pular(2)

    f.linha("Movimentação")
    f.linha("Movimentação > Ações | Aluguel")
    f.linha("Data", "Posição", "Código", "Transação", "Taxa %", "Qtde.",
            "Preço de Referência R$", "Valor Contratado R$")
    f.linha(spec.fim - timedelta(days=4), "Doador", tk(_ALUGUEL_TICKER), "EMPRESTIMO",
            _ALUGUEL_TAXA_ANO, _ALUGUEL_QTDE, _ALUGUEL_PRECO_REF, contratado)
    f.pular(2)

    # Posição > ETF
    f.linha("Posição")
    f.linha("Posição > ETF")
    f.linha("Código", "Ativo", "Qtde.", "Preço Fechamento R$", "Preço Médio R$", "Saldo Bruto R$")
    for ticker, nome, qtde, preco, medio in _ETFS:
        f.linha(tk(ticker), nome, qtde, preco, medio, _r2(qtde * preco))
    f.linha("Total em ETF's R$", None, None, None, None, _soma_etfs())
    f.pular(2)

    # Posição > Fundos Listados
    f.linha("Posição")
    f.linha("Posição > Fundos Listados")
    f.linha("Código", "Ativo", "Tipo", "Qtde.", "Preço Fechamento R$", "Preço Médio R$",
            "Saldo Bruto R$")
    for ticker, nome, tipo, qtde, preco, medio in _FIIS:
        f.linha(tk(ticker), nome, tipo, qtde, preco, medio, _r2(qtde * preco))
    f.linha("Total em Fundos Listados R$", None, None, None, None, None, _soma_fiis())
    f.pular(2)

    # Movimentação > Fundos Listados (rendimentos dos FIIs)
    f.linha("Movimentação")
    f.linha("Movimentação > Fundos Listados")
    f.linha("Data", "Transação", "Código", "Tipo", "Qtde.", "Preço R$", "Valor Bruto R$",
            "Corretagem e Emolumentos R$", "Valor Líquido R$")
    proventos_fiis = [p for p in _PROVENTOS if p[3] == "fiis"]
    for offset, ticker, transacao, _, qtde, bruto, liquido in proventos_fiis:
        f.linha(_dia(spec, offset), transacao, tk(ticker), "FII", qtde, _NULO, bruto, _NULO, liquido)
    f.linha("Total de Compras", None, None, None, None, None, _NULO, _NULO, _NULO)
    f.linha("Total de Vendas", None, None, None, None, None, _NULO, _NULO, _NULO)
    f.linha("Total de Proventos", None, None, None, None, None,
            _r2(sum(p[5] for p in proventos_fiis)), None, _r2(sum(p[6] for p in proventos_fiis)))
    _rodape_decoracao(f, spec)


def _aba_fundos(wb: Workbook, spec: EspecExtrato) -> None:
    f = _Folha(wb, "Fundos")
    anterior = spec.inicio - timedelta(days=1)
    bruto = _r2(_FUNDO_QTDE_COTAS * _FUNDO_COTACAO)
    liquido = _r2(bruto - _FUNDO_IR)
    f.pular()
    f.linha("Fundos")
    f.pular()
    f.linha("Posições")
    f.pular()
    f.linha("Posição > Portfólio de fundos")
    f.pular()
    f.linha("Data Referência", f"Saldo Líquido R$ {anterior:%d/%m/%y}", "Quantidade de Cotas",
            "Cotação Atual R$", "Saldo Bruto R$", "Provisão de IR R$", "Provisão de IOF R$",
            "Saldo Líquido R$", "Variação Nominal R$")
    # Registro em DUAS linhas: o nome do fundo (com a classe CNPJ) numa linha, os dados na
    # seguinte — é a armadilha real da aba Fundos que o parser precisa atravessar.
    f.linha(f"{_FUNDO_NOME} - Classe CNPJ: {_FUNDO_CNPJ}")
    f.linha(spec.fim, _FUNDO_LIQ_ANTERIOR, _FUNDO_QTDE_COTAS, _FUNDO_COTACAO, bruto,
            _FUNDO_IR, _NULO, liquido, _r2(liquido - _FUNDO_LIQ_ANTERIOR))
    f.pular()
    f.linha("Total em fundos", None, None, None, bruto, _NULO, _NULO, liquido, None)
    f.pular()
    f.linha("Movimentações")
    f.pular()
    f.linha(f"Movimentação > {_FUNDO_NOME}")
    f.linha("Data", "Transação", "Quantidade de Cotas", "Valor da Cota R$", "Valor Bruto R$",
            "IR R$", "IOF R$", "Valor Líquido R$")
    cotas, valor_cota = _FUNDO_APLICACAO
    aplicado = _r2(cotas * valor_cota)
    f.linha(spec.inicio, "APLICAÇÃO", cotas, valor_cota, aplicado, _NULO, _NULO, aplicado)
    f.linha("Total de Aplicações", None, cotas, None, aplicado, _NULO, _NULO, aplicado)
    f.linha("Total de Resgates", None, _NULO, None, _NULO, _NULO, _NULO, _NULO)


def _aba_cripto(wb: Workbook, spec: EspecExtrato) -> None:
    f = _Folha(wb, "CriptoAtivos")
    f.pular()
    f.linha("CriptoAtivos")
    f.pular()
    if not spec.cripto_so_movimentacao:
        valor = _r2(_CRIPTO_QTDE * _CRIPTO_PRECO)
        f.linha("Posição")
        f.pular()
        f.linha("Posição > Portfólio de CriptoAtivos")
        # 'Valor Liquido' SEM acento: grafia real desta aba
        f.linha("Ativo", "Quantidade", "Preço R$", "Valor Liquido R$", "IR R$", "IOF R$",
                "Valor Bruto R$")
        f.linha(_CRIPTO_ATIVO, _CRIPTO_QTDE, _CRIPTO_PRECO, valor, _NULO, _NULO, valor)
        f.linha("Total", None, None, None, None, None, valor)
        f.pular(2)
    offset, ativo, descricao, qtde, preco, valor_mov = _CRIPTO_MOV
    f.linha("Movimentações")
    f.pular()
    f.linha("Movimentação > Portfólio de CriptoAtivos")
    f.linha("Data", "Ativo", "Descrição", "Qtde.", "Preço R$", "Valor Liquido R$", "IR R$",
            "IOF R$", "Valor Bruto R$")
    f.linha(_dia(spec, offset), ativo, descricao, qtde, preco, valor_mov, _NULO, _NULO, valor_mov)


def _aba_conta_corrente(wb: Workbook, spec: EspecExtrato) -> None:
    f = _Folha(wb, "Conta Corrente")
    saldo_anterior, linhas, saldo_final = _razao(spec)
    f.pular()
    f.linha("Conta Corrente")
    f.pular()
    f.linha("Posição")
    f.pular()
    f.linha("Data", "Valor financeiro R$")
    f.linha(spec.fim, saldo_final)
    f.pular(2)
    f.linha("Movimentações")
    f.pular()
    f.linha("Data", "Descrição", "Movimentação R$", "Saldo conta investimentos R$")
    f.linha(spec.inicio, "Saldo Anterior", None, saldo_anterior)
    for d, desc, valor, saldo in linhas:
        f.linha(d, desc, valor, saldo)
    creditos = _r2(sum(v for _, _, v, _ in linhas if v > 0))
    debitos = _r2(sum(v for _, _, v, _ in linhas if v < 0))
    f.linha("Total de Créditos", None, creditos if creditos else _NULO)
    f.linha("Total de Débitos", None, debitos if debitos else _NULO)


def _aba_transito(wb: Workbook, spec: EspecExtrato) -> None:
    f = _Folha(wb, "Valores em Trânsito")
    f.pular()
    f.linha("Valores em Trânsito")
    f.pular()
    f.linha("Renda Variavel")
    f.pular()
    f.linha("Data Liquidação", "Descrição", "Valor R$")
    for offset, desc, valor in _TRANSITO:
        f.linha(spec.fim + timedelta(days=offset), desc, valor)
    f.linha("Total", None, _transito_total())


def _aba_fale_conosco(wb: Workbook) -> None:
    f = _Folha(wb, "Fale Conosco")
    f.pular()
    f.linha("Fale Conosco")
    f.pular(2)
    f.linha("Nosso atendimento")
    f.linha("4007-2511 Regiões metropolitanas.")
    f.linha("0800-001-251 Demais localidades.")
    f.linha("24 horas por dia, 7 dias por semana.")
    f.linha("atendimentoinvestimentos@btgpactual.com")


def construir(spec: EspecExtrato | None = None) -> Workbook:
    """Workbook openpyxl do extrato sintético — do zero, nada é lido de arquivo."""
    spec = spec or EspecExtrato()
    wb = Workbook()
    wb.remove(wb.active)
    _aba_capa(wb, spec)
    _aba_sumario(wb, spec)
    if spec.com_fundos:
        _aba_fundos(wb, spec)
    _aba_renda_fixa(wb, spec)
    _aba_renda_variavel(wb, spec)
    if spec.com_cripto:
        _aba_cripto(wb, spec)
    _aba_conta_corrente(wb, spec)
    _aba_transito(wb, spec)
    _aba_fale_conosco(wb)
    return wb


def fixture_padrao() -> Workbook:
    """O extrato sintético padrão — o substituto da antiga fixture de arquivo."""
    return construir(EspecExtrato())


def gerar(spec: EspecExtrato | None = None) -> bytes:
    """O extrato da spec como bytes de XLSX, pronto para o parser ou para upload."""
    wb = construir(spec)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Constantes públicas — derivadas das MESMAS tabelas que o builder escreve.
# Os testes asseram contra elas; nenhum teste deve repetir um número na mão.
# ---------------------------------------------------------------------------

_ESPEC_PADRAO = EspecExtrato()
_SOMAS_SUMARIO = {nome.rstrip("*"): (b_ini, b_fim)
                  for nome, b_ini, _, b_fim, _ in _linhas_sumario(_ESPEC_PADRAO)}

DATA_REFERENCIA_FIXTURE = _ESPEC_PADRAO.fim.isoformat()           # "2026-07-31"
DATA_ANTERIOR_FIXTURE = (_ESPEC_PADRAO.inicio - timedelta(days=1)).isoformat()

TOTAL_POSICOES_FIXTURE = len(_ACOES) + len(_ETFS) + len(_FIIS) + len(_posicoes_td()) + 1
LOTES_RF_FIXTURE = len(_LOTES_TD)                                 # 8
LANCAMENTOS_FIXTURE = len(_PROVENTOS) + 1                         # + rendimento do saldo
PROVENTOS_FIXTURE = len(_PROVENTOS)                               # 4
ITENS_TRANSITO_FIXTURE = len(_TRANSITO)                           # 3

TOTAL_BRUTO_INI_FIXTURE = _SOMAS_SUMARIO["Total"][0]
TOTAL_BRUTO_FIM_FIXTURE = _SOMAS_SUMARIO["Total"][1]
RV_SUMARIO_FIM_FIXTURE = _SOMAS_SUMARIO["Renda Variável"][1]
RF_BRUTO_FIXTURE = _rf_totais()[0]
RF_LIQUIDO_FIXTURE = _rf_totais()[1]
RV_POSICOES_FIXTURE = _r2(_soma_acoes() + _soma_etfs() + _soma_fiis())
CAIXA_INI_FIXTURE = _CAIXA_INI
CAIXA_FIM_FIXTURE = _razao(_ESPEC_PADRAO)[2]
TRANSITO_FIXTURE = _transito_total()
# Soma de TODAS as posições parseáveis (RV + RF + caixa) — o total_parseado da checagem.
TOTAL_POSICOES_VALOR_FIXTURE = _r2(RV_POSICOES_FIXTURE + RF_BRUTO_FIXTURE + CAIXA_FIM_FIXTURE)

VARIACAO_MES_FIXTURE = _r2(TOTAL_BRUTO_FIM_FIXTURE - TOTAL_BRUTO_INI_FIXTURE)
RENTABILIDADE_PCT_FIXTURE = round(VARIACAO_MES_FIXTURE / TOTAL_BRUTO_INI_FIXTURE * 100, 4)

PROVENTOS_LIQUIDOS_FIXTURE = _r2(sum(p[6] for p in _PROVENTOS))
PROVENTOS_POR_CLASSE_FIXTURE = {
    "ACAO": _r2(sum(p[6] for p in _PROVENTOS if p[3] == "acoes")),
    "FII": _r2(sum(p[6] for p in _PROVENTOS if p[3] == "fiis")),
}
RENDIMENTO_SALDO_FIXTURE = _RENDIMENTO_SALDO
RENDIMENTO_KNCR11_FIXTURE = next(p[6] for p in _PROVENTOS if p[1] == "KNCR11")
CREDITOS_FIXTURE = _r2(PROVENTOS_LIQUIDOS_FIXTURE + RENDIMENTO_SALDO_FIXTURE)
# Valores do razão na ordem das linhas (4 proventos + rendimento do saldo remunerado)
VALORES_RAZAO_FIXTURE = [p[6] for p in _PROVENTOS] + [_RENDIMENTO_SALDO]

ALUGUEL_RESULTADO_FIXTURE = _ESPEC_PADRAO.aluguel_resultado
ALUGUEL_CONTRATADO_FIXTURE = _r2(_ALUGUEL_QTDE * _ALUGUEL_PRECO_REF)
TAEE11_QTDE_FIXTURE = next(q for t, _, q, _, _ in _ACOES if t == "TAEE11")
TAEE11_VALOR_FIXTURE = _r2(next(q * p for t, _, q, p, _ in _ACOES if t == "TAEE11"))

# Tesouro IPCA+ 2029 (NTNB-P): custo real de aquisição vem dos 4 lotes do Detalhamento.
_NTNB = [l for l in _LOTES_TD if l.sigla == "NTNB-P"]
NTNB_QTDE_FIXTURE = _r2(sum(l.qtde for l in _NTNB))
NTNB_CUSTO_TOTAL_FIXTURE = _r2(sum(l.valor_compra for l in _NTNB))
NTNB_PRECO_MEDIO_FIXTURE = _r2(NTNB_CUSTO_TOTAL_FIXTURE / NTNB_QTDE_FIXTURE)
NTNB_PRECO_ATUAL_FIXTURE = _NTNB[0].preco_atual
NTNB_LOTES_FIXTURE = len(_NTNB)
NTNB_VENCIMENTO_FIXTURE = _NTNB[0].vencimento.isoformat()
NTNB_TAXA_FIXTURE = _TAXA_MEDIA_TD[("NTNB-P", NTNB_VENCIMENTO_FIXTURE)]

# LFT 2031 — lote único: o custo é o preço de compra declarado, sem divisão.
_LFT31 = next(l for l in _LOTES_TD if l.sigla == "LFT" and l.vencimento.year == 2031)
LFT31_PRECO_COMPRA_FIXTURE = _LFT31.preco_compra
LFT31_VALOR_COMPRA_FIXTURE = _LFT31.valor_compra
LFT31_PRECO_ATUAL_FIXTURE = _LFT31.preco_atual
LFT31_AQUISICAO_FIXTURE = _LFT31.aquisicao.isoformat()
LFT31_TAXA_FIXTURE = _LFT31.taxa
