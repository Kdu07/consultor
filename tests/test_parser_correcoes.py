"""
Correções v3 do parser (plano "Confiabilidade dos extratos", WS-A):

  - _limpar_codigo: tickers/siglas decorados com '*' saem limpos (chave estável) e avisados;
  - _tabelas v2: nada se perde em silêncio (linha órfã numérica → linhas_ignoradas;
    linha com cara de cabeçalho no meio de tabela aberta é dado);
  - abas Fundos (registro em duas linhas, fundo encerrado) e CriptoAtivos;
  - quantidade/vencimento ilegíveis viram aviso tipado, nunca descarte mudo;
  - checagem desconta o resultado do aluguel que o Sumário do BTG soma ao RV;
  - payload v3: totais_abas + avisos_parser.

Os XLSX são construídos inline (openpyxl) com VALORES INVENTADOS — nenhum número de
carteira real entra em arquivo versionado. Não usa tests/planilhas.py de propósito:
a fixture está sendo migrada em paralelo e este arquivo não pode depender dela.
"""
from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.tools.btg_xlsx_parser import (
    VERSAO_PARSER,
    _limpar_codigo,
    _tabelas,
    parse_btg_xlsx,
)

# ---------------------------------------------------------------------------
# Universo sintético (valores inventados, coerentes entre si)
# ---------------------------------------------------------------------------

RV_ACOES = 2550.0          # BBAS3: 100 × 25,50
RV_ETF = 300.0             # IVVB11: 10 × 30,00
ALUGUEL_RESULT = 1.5       # resultado acumulado do aluguel (o Sumário soma isto ao RV)
RV_SUMARIO = RV_ACOES + RV_ETF + ALUGUEL_RESULT
RF_BRUTO, RF_LIQ = 24000.0, 23900.0
FUNDO_BRUTO, FUNDO_LIQ = 1050.0, 1040.0
CRIPTO_BRUTO = 5000.0
CAIXA = 500.0
TRANSITO = 30.0
TOTAL_BRUTO = RV_SUMARIO + RF_BRUTO + FUNDO_BRUTO + CRIPTO_BRUTO + CAIXA + TRANSITO
TOTAL_LIQ = RV_SUMARIO + RF_LIQ + FUNDO_LIQ + CRIPTO_BRUTO + CAIXA + TRANSITO

ANT = {  # período anterior (31/05/26): só o Sumário precisa fechar
    "rv": (2800.0, 2790.0), "rf": (23800.0, 23700.0), "fundos": (1000.0, 995.0),
    "cripto": (4900.0, 4900.0), "cc": (100.0, 100.0), "transito": (0.0, 0.0),
}
ANT_TOTAL_BRUTO = sum(v[0] for v in ANT.values())
ANT_TOTAL_LIQ = sum(v[1] for v in ANT.values())

FUNDO_NOME = "Fundo Teste Alfa FIC FIM"
FUNDO_CNPJ = "11.222.333/0001-44"
FUNDO_CHAVE = "FUNDO:11-222-333-0001-44"


def montar_xlsx(abas: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for nome, linhas in abas.items():
        ws = wb.create_sheet(nome)
        for linha in linhas:
            ws.append(linha)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def aba_capa() -> list[list]:
    return [["Período de 01/06/26 a 30/06/26"], ["Emitido em 01/07/26"]]


def aba_sumario(decorar: bool = False) -> list[list]:
    # Período aberto decora os NOMES dos mercados ('Renda Variável*') — visto em extrato
    # real; o parser tem de casar mesmo assim.
    estrela = "*" if decorar else ""
    return [
        ["Sumário"],
        ["Mercados", "Saldo Bruto R$ 31/05/26", "Saldo Líquido R$ 31/05/26",
         "Saldo Bruto R$ 30/06/26", "Saldo Líquido R$ 30/06/26"],
        [f"Renda Variável{estrela}", *ANT["rv"], RV_SUMARIO, RV_SUMARIO],
        [f"Renda Fixa{estrela}", *ANT["rf"], RF_BRUTO, RF_LIQ],
        ["Fundos de Investimento", *ANT["fundos"], FUNDO_BRUTO, FUNDO_LIQ],
        ["CriptoAtivos", *ANT["cripto"], CRIPTO_BRUTO, CRIPTO_BRUTO],
        ["Conta Corrente", *ANT["cc"], CAIXA, CAIXA],
        ["Valores em Trânsito", *ANT["transito"], TRANSITO, TRANSITO],
        ["Total", ANT_TOTAL_BRUTO, ANT_TOTAL_LIQ, TOTAL_BRUTO, TOTAL_LIQ],
    ]


def aba_rv(codigo_acao: str = "BBAS3", qtde_acao=100.0) -> list[list]:
    return [
        ["Renda Variável"],
        ["Posição > Ações"],
        ["Código", "Ação", "Qtde.", "Preço Fechamento R$", "Preço Médio R$", "Saldo Bruto R$"],
        [codigo_acao, "BRASIL ON NM", qtde_acao, 25.50, 20.00, RV_ACOES],
        ["Total em Ações R$", "", "", "", "", RV_ACOES],
        [],
        ["Posição > ETF"],
        ["Código", "Ativo", "Qtde.", "Preço Fechamento R$", "Preço Médio R$", "Saldo Bruto R$"],
        ["IVVB11", "ISHARE SP500", 10.0, 30.00, 28.00, RV_ETF],
        ["Total em ETF's R$", "", "", "", "", RV_ETF],
        [],
        ["Posição > Ações | Aluguel"],
        ["Código", "Qtde.", "Posição", "Valor Contratado R$", "Data Vencimento",
         "Taxa Ano %", "Result. Acum. Líq. R$"],
        ["TAEE11*", 10.0, "Doador", 350.0, "21/10/26", 1.5, ALUGUEL_RESULT],
        ["Total em Aluguel de Ações R$", "", "", 350.0, "", "", ALUGUEL_RESULT],
        [],
        ["Movimentação > Ações"],
        ["Data", "Transação", "Código", "Qtde.", "Preço R$", "Valor Bruto R$",
         "Corretagem e Emolumentos R$", "Valor Líquido R$"],
        ["05/06/26", "COMPRA À VISTA", "BBAS3*", 10.0, 25.0, 250.0, 0.5, 250.5],
        ["10/06/26", "RECEBIMENTO DIVIDENDOS", "BBAS3*", "-", "-", 12.0, "-", 12.0],
        ["Total de Compras", "", "", "", "", 250.0, 0.5, 250.5],
        ["Total de Proventos", "", "", "", "", 12.0, "", 12.0],
    ]


def aba_rf(vencimento="01/03/31", segunda_linha: list | None = None) -> list[list]:
    linhas = [
        ["Renda Fixa"],
        ["Posição > TESOURO DIRETO - LFT"],
        ["Emissor", "Ativo", "Vencimento", "Taxa Média Ponderada", "Quantidade",
         "Preço R$", "Saldo Bruto R$", "IR R$", "IOF R$", "Saldo Líquido R$"],
        ["BACEN-BANCO CENTRAL DO BRASIL - RJ", "LFT*", vencimento, "SELIC + 0,04%",
         1.5, 16000.0, RF_BRUTO, 100.0, "-", RF_LIQ],
    ]
    if segunda_linha is not None:
        linhas.append(segunda_linha)
    linhas.append(["Total", "", "", "", "", "", RF_BRUTO, 100.0, "-", RF_LIQ])
    return linhas


def aba_cc() -> list[list]:
    return [
        ["Conta Corrente"],
        ["Posição"],
        ["Data", "Valor financeiro R$"],
        ["30/06/26", CAIXA],
        [],
        ["Movimentações"],
        ["Data", "Descrição", "Movimentação R$", "Saldo conta investimentos R$"],
        ["01/06/26", "Saldo Anterior", "", 100.0],
        ["05/06/26", "PIX RECEBIDO FULANO DE TAL", 400.0, 500.0],
        ["Total de Créditos", "", 400.0, ""],
        ["Total de Débitos", "", "-", ""],
    ]


def aba_transito() -> list[list]:
    return [
        ["Valores em Trânsito"],
        ["Renda Variavel"],
        ["Data Liquidação", "Descrição", "Valor R$"],
        ["14/07/26", "RENDIMENTO - FII TESTE", TRANSITO],
        ["Total", "", TRANSITO],
    ]


def aba_fundos(encerrado: bool = False) -> list[list]:
    # Registro em DUAS linhas (layout real): nome numa célula só, dados na seguinte.
    # Encerrado: dados atuais '-' e só o Saldo Líquido do período anterior preenchido.
    dados = (
        ["-", 1000.0, "-", "-", "-", "-", "-", "-", 40.0]
        if encerrado else
        ["30/06/26", 1000.0, 10.0, 105.0, FUNDO_BRUTO, 10.0, "-", FUNDO_LIQ, 40.0]
    )
    total = (
        ["Total em fundos", "", "", "", "-", "-", "-", "-", 40.0]
        if encerrado else
        ["Total em fundos", "", "", "", FUNDO_BRUTO, 10.0, "-", FUNDO_LIQ, 40.0]
    )
    return [
        ["Fundos"],
        ["Posições"],
        ["Posição > Portfólio de fundos"],
        ["Data Referência", "Saldo Líquido R$ 31/05/26", "Quantidade de Cotas",
         "Cotação Atual R$", "Saldo Bruto R$", "Provisão de IR R$", "Provisão de IOF R$",
         "Saldo Líquido R$", "Variação Nominal R$"],
        [f"{FUNDO_NOME} - Classe CNPJ: {FUNDO_CNPJ}"],
        dados,
        total,
        [],
        ["Movimentações"],
        [f"Movimentação > {FUNDO_NOME}"],
        ["Data", "Transação", "Quantidade de Cotas", "Valor da Cota R$", "Valor Bruto R$",
         "IR R$", "IOF R$", "Valor Líquido R$"],
        ["10/06/26", "APLICAÇÃO", 2.0, 100.0, 200.0, "-", "-", 200.0],
        ["Total de Aplicações", "", 2.0, "", 200.0, "-", "-", 200.0],
        ["Total de Resgates", "", "-", "", "-", "-", "-", "-"],
    ]


def aba_cripto(so_movimentacao: bool = False) -> list[list]:
    # 'Valor Liquido' SEM acento é o rótulo real desta aba.
    posicao = [
        ["Posição"],
        ["Posição > Portfólio de CriptoAtivos"],
        ["Ativo", "Quantidade", "Preço R$", "Valor Liquido R$", "IR R$", "IOF R$", "Valor Bruto R$"],
        ["XBT BITCOIN (XBT)", 0.01, 500000.0, CRIPTO_BRUTO, "-", "-", CRIPTO_BRUTO],
        ["Total", "", "", CRIPTO_BRUTO, "", "", CRIPTO_BRUTO],
        [],
    ]
    return [
        ["CriptoAtivos"],
        *([] if so_movimentacao else posicao),
        ["Movimentações"],
        ["Movimentação > Portfólio de CriptoAtivos"],
        ["Data", "Ativo", "Descrição", "Qtde.", "Preço R$", "Valor Liquido R$",
         "IR R$", "IOF R$", "Valor Bruto R$"],
        ["19/06/26", "XBT BITCOIN (XBT)", "Venda", 0.005, 500000.0, 2500.0, "-", "-", 2500.0],
    ]


def xlsx_completo(decorar_sumario: bool = False) -> bytes:
    """Extrato sintético com todas as abas, fechando ao centavo (veredito 'ok')."""
    return montar_xlsx({
        "Capa": aba_capa(),
        "Sumario": aba_sumario(decorar=decorar_sumario),
        "Fundos": aba_fundos(),
        "Renda Fixa": aba_rf(),
        "Renda Variavel": aba_rv(),
        "CriptoAtivos": aba_cripto(),
        "Conta Corrente": aba_cc(),
        "Valores em Trânsito": aba_transito(),
    })


def _por_chave(extrato, chave):
    return next((p for p in extrato.posicoes if p.chave_externa == chave), None)


def _tipos(extrato) -> set[str]:
    return {a["tipo"] for a in extrato.avisos_parser}


# ---------------------------------------------------------------------------
# _limpar_codigo
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("bruto, esperado, decorado", [
    ("BBAS3* ", "BBAS3", True),
    ("BBAS3*", "BBAS3", True),
    ("LFT* ", "LFT", True),
    ("LFT *", "LFT", True),
    ("BBAS3", "BBAS3", False),
    ("bbas3", "BBAS3", False),       # caixa baixa não é decoração
    ("NTNB-P", "NTNB-P", False),     # hífen de sigla sobrevive
    (None, "", False),
    ("", "", False),
])
def test_limpar_codigo(bruto, esperado, decorado):
    assert _limpar_codigo(bruto) == (esperado, decorado)


def test_rv_ticker_decorado_sai_limpo_com_aviso():
    extrato = parse_btg_xlsx(montar_xlsx({
        "Capa": aba_capa(), "Renda Variavel": aba_rv(codigo_acao="BBAS3*"),
    }))
    pos = _por_chave(extrato, "B3:BBAS3")
    assert pos is not None and pos.ticker == "BBAS3"
    assert "ticker_decorado" in _tipos(extrato)
    # movimentação e aluguel também saem limpos
    assert {m["ticker"] for m in extrato.movimentacoes} == {"BBAS3"}
    assert extrato.aluguel[0]["ticker"] == "TAEE11"


def test_rf_sigla_decorada_acha_nome_do_tesouro():
    # 'LFT*' tem de achar 'Tesouro Selic' em _TESOURO_NOME e gerar chave limpa.
    extrato = parse_btg_xlsx(montar_xlsx({"Capa": aba_capa(), "Renda Fixa": aba_rf()}))
    pos = _por_chave(extrato, "TD:LFT:2031-03-01")
    assert pos is not None
    assert pos.nome.startswith("Tesouro Selic")
    assert "ticker_decorado" in _tipos(extrato)


def test_ticker_fora_do_padrao_b3_vira_aviso_nunca_descarte():
    extrato = parse_btg_xlsx(montar_xlsx({
        "Capa": aba_capa(), "Renda Variavel": aba_rv(codigo_acao="ABCDE11"),
    }))
    assert _por_chave(extrato, "B3:ABCDE11") is not None
    assert "ticker_fora_do_padrao_b3" in _tipos(extrato)


# ---------------------------------------------------------------------------
# _tabelas v2 — nada se perde em silêncio
# ---------------------------------------------------------------------------

def _ws(linhas):
    wb = Workbook()
    ws = wb.active
    for linha in linhas:
        ws.append(linha)
    return ws


def test_tabelas_linha_orfa_numerica_vai_para_ignoradas():
    ignoradas, avisos = [], []
    tabelas = _tabelas(_ws([
        ["Posição > Ações"],
        ["Código", "Saldo Bruto R$"],
        ["AAAA3", 10.0],
        [],
        [None, 123.45],          # órfã com número: pode ser dado — registra
    ]), "RV", ignoradas, avisos)
    assert len(tabelas) == 1 and len(tabelas[0].linhas) == 1
    assert len(ignoradas) == 1
    assert "RV" in ignoradas[0] and "123.45" in ignoradas[0]


def test_tabelas_linha_de_uma_celula_no_meio_do_bloco_nao_some_em_silencio():
    # A linha de 1 célula vira "título" e fecha a tabela (comportamento herdado); o que
    # NÃO pode acontecer é a linha de dados seguinte sumir sem rastro.
    ignoradas, avisos = [], []
    tabelas = _tabelas(_ws([
        ["Posição > Ações"],
        ["Código", "Saldo Bruto R$"],
        ["AAAA3", 10.0],
        ["Observação solta no meio do bloco"],
        ["BBBB4", 20.0],
    ]), "RV", ignoradas, avisos)
    assert len(tabelas[0].linhas) == 1
    assert len(ignoradas) == 1 and "BBBB4" in ignoradas[0]


def test_tabelas_linha_cheia_de_tracos_vira_dado_nao_cabecalho():
    # Linha com '-' célula a célula é texto puro — indistinguível de cabeçalho. Com
    # tabela aberta e sem totais, tem de virar DADO (fundo encerrado, cripto zerada).
    tabelas = _tabelas(_ws([
        ["Posição > Portfólio de CriptoAtivos"],
        ["Ativo", "Quantidade", "Valor Bruto R$"],
        ["XBT BITCOIN (XBT)", 0.01, 500.0],
        ["ETH ETHEREUM (ETH)", "-", "-"],
    ]))
    assert len(tabelas) == 1
    assert len(tabelas[0].linhas) == 2


def test_tabelas_cabecalho_depois_de_total_abre_tabela_nova():
    tabelas = _tabelas(_ws([
        ["Posição > Ações"],
        ["Código", "Saldo Bruto R$"],
        ["AAAA3", 10.0],
        ["Total", 10.0],
        ["Código", "Qtde."],       # tabela anterior já fechou com total → abre nova
        ["BBBB4", 5.0],
    ]))
    assert len(tabelas) == 2
    assert len(tabelas[0].linhas) == 1 and len(tabelas[1].linhas) == 1


def test_tabelas_linha_so_texto_fora_de_tabela_vira_contador():
    ignoradas, avisos = [], []
    _tabelas(_ws([
        ["Posição > Ações"],
        ["Código", "Saldo Bruto R$"],
        ["AAAA3", 10.0],
        [],
        [datetime(2026, 6, 1), "nota decorativa sem número"],
    ]), "RV", ignoradas, avisos)
    assert ignoradas == []
    assert any(a["tipo"] == "linhas_texto_fora_de_tabela" for a in avisos)


# ---------------------------------------------------------------------------
# Quantidade e vencimento ilegíveis
# ---------------------------------------------------------------------------

def test_quantidade_ilegivel_derivada_de_saldo_e_preco():
    extrato = parse_btg_xlsx(montar_xlsx({
        "Capa": aba_capa(), "Renda Variavel": aba_rv(qtde_acao="-"),
    }))
    pos = _por_chave(extrato, "B3:BBAS3")
    assert pos.quantidade == pytest.approx(RV_ACOES / 25.50)
    assert "quantidade_derivada" in _tipos(extrato)


def test_quantidade_ilegivel_sem_preco_fica_zero_com_aviso():
    abas = {"Capa": aba_capa(), "Renda Variavel": [
        ["Renda Variável"],
        ["Posição > Ações"],
        ["Código", "Ação", "Qtde.", "Preço Fechamento R$", "Preço Médio R$", "Saldo Bruto R$"],
        ["BBAS3", "BRASIL ON NM", "-", "-", 20.0, RV_ACOES],
    ]}
    extrato = parse_btg_xlsx(montar_xlsx(abas))
    pos = _por_chave(extrato, "B3:BBAS3")
    assert pos.quantidade == 0.0
    assert "quantidade_ilegivel" in _tipos(extrato)


def test_vencimento_ilegivel_avisa_e_colisao_tambem():
    segunda = ["BACEN-BANCO CENTRAL DO BRASIL - RJ", "LFT", "data inválida", "SELIC",
               1.0, 16000.0, 16000.0, "-", "-", 16000.0]
    extrato = parse_btg_xlsx(montar_xlsx({
        "Capa": aba_capa(),
        "Renda Fixa": aba_rf(vencimento="também inválida", segunda_linha=segunda),
    }))
    tipos = _tipos(extrato)
    assert "vencimento_ilegivel" in tipos
    # dois títulos caíram na MESMA chave 'sem-vencimento' → aviso que o validador
    # promove a erro (a posição agregada seria mentira)
    assert "colisao_sem_vencimento" in tipos
    assert _por_chave(extrato, "TD:LFT:sem-vencimento") is not None


# ---------------------------------------------------------------------------
# Fundos e CriptoAtivos
# ---------------------------------------------------------------------------

def test_fundos_posicao_ativa_registro_em_duas_linhas():
    extrato = parse_btg_xlsx(montar_xlsx({"Capa": aba_capa(), "Fundos": aba_fundos()}))
    pos = _por_chave(extrato, FUNDO_CHAVE)
    assert pos is not None
    assert pos.classe == "FUNDO"
    assert pos.nome == FUNDO_NOME                 # sem o sufixo '- Classe CNPJ: ...'
    assert pos.quantidade == 10.0
    assert pos.valor_mercado == FUNDO_BRUTO       # bruto, como nas demais classes
    assert pos.preco_fechamento == 105.0
    assert pos.as_of == "2026-06-30"


def test_fundos_movimentacao_vira_fluxo_com_classe_fundo():
    extrato = parse_btg_xlsx(montar_xlsx({"Capa": aba_capa(), "Fundos": aba_fundos()}))
    movs = [m for m in extrato.movimentacoes if m["classe"] == "FUNDO"]
    assert len(movs) == 1
    assert movs[0]["operacao"] == "COMPRA"        # APLICAÇÃO ≈ compra
    assert movs[0]["chave_externa"] == FUNDO_CHAVE
    assert movs[0]["ticker"] is None              # nada de chave B3 para fundo
    assert movs[0]["valor_liquido"] == 200.0


def test_fundo_encerrado_nao_vira_posicao_mas_avisa():
    extrato = parse_btg_xlsx(montar_xlsx({
        "Capa": aba_capa(), "Fundos": aba_fundos(encerrado=True), "Conta Corrente": aba_cc(),
    }))
    assert not [p for p in extrato.posicoes if p.classe == "FUNDO"]
    assert "fundo_encerrado" in _tipos(extrato)
    # a linha de total ('-' = zero) foi capturada mesmo assim
    totais = extrato.totais_abas["fundos"]
    assert any(t["rotulo"].startswith("Total em fundos") for t in totais)


def test_cripto_posicao_e_movimentacao():
    extrato = parse_btg_xlsx(montar_xlsx({"Capa": aba_capa(), "CriptoAtivos": aba_cripto()}))
    pos = _por_chave(extrato, "CRIPTO:XBT")
    assert pos is not None
    assert pos.classe == "CRIPTO" and pos.ticker == "XBT"
    assert pos.quantidade == 0.01 and pos.valor_mercado == CRIPTO_BRUTO
    movs = [m for m in extrato.movimentacoes if m["classe"] == "CRIPTO"]
    assert movs and movs[0]["operacao"] == "VENDA"
    assert movs[0]["chave_externa"] == "CRIPTO:XBT"


def test_cripto_so_movimentacao_nao_inventa_posicao():
    extrato = parse_btg_xlsx(montar_xlsx({
        "Capa": aba_capa(),
        "CriptoAtivos": aba_cripto(so_movimentacao=True),
        "Conta Corrente": aba_cc(),
    }))
    assert not [p for p in extrato.posicoes if p.classe == "CRIPTO"]
    assert [m for m in extrato.movimentacoes if m["classe"] == "CRIPTO"]


# ---------------------------------------------------------------------------
# Checagem e payload v3
# ---------------------------------------------------------------------------

def test_checagem_desconta_resultado_do_aluguel():
    # O Sumário soma o resultado do aluguel ao RV; sem o desconto este extrato acusaria
    # diferença permanente de ALUGUEL_RESULT (o falso alarme antigo).
    extrato = parse_btg_xlsx(xlsx_completo())
    assert extrato.checagem["ok"] is True
    assert extrato.checagem["rendimento_aluguel_excluido"] == ALUGUEL_RESULT
    por_mercado = {c["mercado"]: c["diferenca"] for c in extrato.checagem["por_mercado"]}
    assert por_mercado["Renda Variável"] == 0.0
    assert por_mercado["Fundos de Investimento"] == 0.0
    assert por_mercado["CriptoAtivos"] == 0.0
    assert extrato.checagem["diferenca"] == 0.0


def test_checagem_casa_mercado_decorado_com_estrela():
    extrato = parse_btg_xlsx(xlsx_completo(decorar_sumario=True))
    mercados = {c["mercado"] for c in extrato.checagem["por_mercado"]}
    assert "Renda Variável" in mercados
    assert extrato.checagem["ok"] is True


def test_payload_v3_tem_totais_abas_e_avisos():
    extrato = parse_btg_xlsx(xlsx_completo())
    assert extrato.versao_parser == VERSAO_PARSER == 3
    payload = extrato.to_dict()
    assert set(payload["totais_abas"]) == {
        "renda_variavel", "renda_fixa", "fundos", "cripto", "valores_em_transito",
    }
    acoes = next(t for t in payload["totais_abas"]["renda_variavel"]
                 if t["bloco"] == "Posição > Ações")
    assert acoes["valores"]["saldo bruto r$"] == RV_ACOES
    assert acoes["soma_linhas"]["saldo bruto r$"] == RV_ACOES
    assert isinstance(payload["avisos_parser"], list)
    # decoração do aluguel sintético ('TAEE11*') registrada
    assert any(a["tipo"] == "ticker_decorado" for a in payload["avisos_parser"])
