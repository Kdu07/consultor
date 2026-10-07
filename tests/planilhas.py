"""
Alterações sobre o extrato sintético padrão (tests/gerador_extrato.py).

A fixture não é mais um arquivo: abrir_fixture() devolve o workbook que o gerador monta do
zero, com a forma do extrato real do BTG e valores inventados. Os helpers abaixo alteram
esse workbook por rótulo/cabeçalho — aportes, débitos, compras, lotes novos de Tesouro,
outros meses — recalculando o que o parser confere (saldos do razão, totais, posição CAIXA
e Sumário), para que o extrato continue coerente e só o aspecto sob teste mude.

Tudo em memória: abrir_fixture() → alterar → para_bytes(). Nenhum extrato real entra no
repositório; os textos imitam o formato do BTG, nenhum vem de extrato de verdade.

O nome não começa com test_: o pytest não coleta este módulo.
"""
from __future__ import annotations

from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl.workbook import Workbook

from app.tools.btg_xlsx_parser import _norm
from tests.gerador_extrato import fixture_padrao


def abrir_fixture() -> Workbook:
    """O extrato sintético padrão, pronto para os helpers deste módulo alterarem."""
    return fixture_padrao()


def bytes_fixture() -> bytes:
    """O extrato sintético padrão como bytes de XLSX (o que era FIXTURE.read_bytes())."""
    return para_bytes(fixture_padrao())


def para_bytes(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Payload de mês arquivado, sem XLSX — para os testes do motor de desempenho
# ---------------------------------------------------------------------------

def payload_mes(
    ref: str,
    v_ini: float,
    v_fim: float,
    lancamentos: list[tuple[str, str, float]] = (),
    posicoes: list[dict] = (),
    proventos: list[dict] = (),
    versao: int = 2,
    razao_ok: bool | None = True,
    periodo_inicio: str | None = None,
) -> dict:
    """
    O dict que ExtratoImportado.payload_json guardaria para um mês: Sumário (V_ini, V_fim),
    razão da conta (data, descrição, valor) e o que mais o teste precisar. `versao=1`
    imita um arquivo anterior ao parser v2 (sem razão).
    """
    fim = date.fromisoformat(ref)
    inicio = date.fromisoformat(periodo_inicio) if periodo_inicio else fim.replace(day=1)
    anterior = inicio - timedelta(days=1)
    payload = {
        "data_referencia": ref,
        "sumario": {
            "atual": {"data": ref, "mercados": {}, "total": {"bruto": v_fim, "liquido": v_fim}},
            "anterior": {"data": anterior.isoformat(), "mercados": {},
                         "total": {"bruto": v_ini, "liquido": v_ini}},
            "meta": {"periodo_inicio": inicio.isoformat(), "emitido_em": None},
        },
        "posicoes": [dict(p) for p in posicoes],
        "proventos": [dict(p) for p in proventos],
        "movimentacoes": [],
        "checagem": {"ok": True},
    }
    if versao >= 2:
        payload["versao_parser"] = versao
        payload["lancamentos_conta"] = [
            {"seq": i, "data": d, "descricao": desc, "valor": v, "saldo": None}
            for i, (d, desc, v) in enumerate(lancamentos)
        ]
        payload["lotes_rf"] = []
        payload["checagem"]["conta_corrente"] = {"ok": razao_ok}
    return payload


# ---------------------------------------------------------------------------
# Localização
# ---------------------------------------------------------------------------

def _aba(wb: Workbook, nome: str):
    for ws in wb.worksheets:
        if _norm(ws.title) == _norm(nome):
            return ws
    raise KeyError(nome)


def _linha(ws, texto: str, a_partir_de: int = 1, coluna: int | None = None) -> int:
    """Primeira linha (1-based) com célula de texto que começa com `texto` (normalizado)."""
    alvo = _norm(texto)
    colunas = [coluna] if coluna else range(1, ws.max_column + 1)
    for r in range(a_partir_de, ws.max_row + 1):
        for c in colunas:
            v = ws.cell(r, c).value
            if isinstance(v, str) and _norm(v).startswith(alvo):
                return r
    raise LookupError(texto)


def _colunas(ws, linha_cabecalho: int) -> dict[str, int]:
    """Cabeçalho normalizado → índice da coluna."""
    out: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(linha_cabecalho, c).value
        if isinstance(v, str) and v.strip():
            out[_norm(v)] = c
    return out


def _col(colunas: dict[str, int], prefixo: str) -> int:
    alvo = _norm(prefixo)
    if alvo in colunas:
        return colunas[alvo]
    for nome, c in colunas.items():
        if nome.startswith(alvo):
            return c
    raise KeyError(prefixo)


def _dt(d: date | datetime) -> datetime:
    return d if isinstance(d, datetime) else datetime(d.year, d.month, d.day)


# ---------------------------------------------------------------------------
# Sumário
# ---------------------------------------------------------------------------

def _colunas_sumario(ws) -> tuple[int, list[int]]:
    """(linha do cabeçalho, colunas 'Saldo ...' na ordem: bruto/líq anterior, bruto/líq atual)."""
    r = _linha(ws, "Mercados")
    cols = [c for c in range(1, ws.max_column + 1)
            if isinstance(ws.cell(r, c).value, str) and _norm(ws.cell(r, c).value).startswith("saldo")]
    return r, cols


def _somar_no_sumario(wb: Workbook, mercado: str, delta: float, periodo: str = "atual") -> None:
    """Soma `delta` ao mercado e ao Total do Sumário, bruto e líquido."""
    ws = _aba(wb, "Sumario")
    r_cab, cols = _colunas_sumario(ws)
    alvo = cols[2:4] if periodo == "atual" else cols[0:2]
    for rotulo in (mercado, "Total"):
        r = _linha(ws, rotulo, a_partir_de=r_cab + 1)
        for c in alvo:
            ws.cell(r, c).value = round(float(ws.cell(r, c).value or 0) + delta, 2)


def definir_total_anterior(wb: Workbook, total_bruto: float) -> None:
    """Patrimônio do fim do mês anterior (V_ini). Ajusta o líquido pelo mesmo delta."""
    ws = _aba(wb, "Sumario")
    r_cab, cols = _colunas_sumario(ws)
    r = _linha(ws, "Total", a_partir_de=r_cab + 1)
    delta = total_bruto - float(ws.cell(r, cols[0]).value)
    ws.cell(r, cols[0]).value = round(total_bruto, 2)
    ws.cell(r, cols[1]).value = round(float(ws.cell(r, cols[1]).value) + delta, 2)


# ---------------------------------------------------------------------------
# Período
# ---------------------------------------------------------------------------

def definir_periodo(wb: Workbook, inicio: date, fim: date) -> None:
    """
    Muda o período do extrato: linha da Capa, datas das colunas do Sumário (anterior =
    véspera do início) e as datas de saldo da conta corrente. As datas dos lançamentos e
    proventos de julho NÃO são movidas — quem precisa de datas no período as acrescenta.
    """
    capa = _aba(wb, "Capa")
    r = _linha(capa, "Periodo de")
    for c in range(1, capa.max_column + 1):
        v = capa.cell(r, c).value
        if isinstance(v, str) and _norm(v).startswith("periodo de"):
            capa.cell(r, c).value = f"Período de {inicio:%d/%m/%y} a {fim:%d/%m/%y}"

    ws = _aba(wb, "Sumario")
    r_cab, cols = _colunas_sumario(ws)
    vespera = inicio - timedelta(days=1)
    rotulos = [
        f"Saldo Bruto R$ {vespera:%d/%m/%y}", f"Saldo Líquido R$ {vespera:%d/%m/%y}",
        f"Saldo Bruto R$ {fim:%d/%m/%y}", f"Saldo Líquido R$ {fim:%d/%m/%y}",
    ]
    for c, rotulo in zip(cols, rotulos):
        ws.cell(r_cab, c).value = rotulo

    cc = _aba(wb, "Conta Corrente")
    r_pos = _linha(cc, "Data")
    cc.cell(r_pos + 1, 2).value = _dt(fim)
    r_ini = _linha(cc, "Saldo Anterior")
    cc.cell(r_ini, 2).value = _dt(inicio)
    r_fim = _linha(cc, "Saldo Final")
    cc.cell(r_fim, 2).value = _dt(fim)


# ---------------------------------------------------------------------------
# Razão da conta corrente
# ---------------------------------------------------------------------------

def adicionar_lancamentos(
    wb: Workbook,
    lancamentos: list[tuple[date, str, float]],
    debitos_positivos: bool = False,
) -> float:
    """
    Acrescenta lançamentos (data, descrição, valor com sinal) ao razão, em ordem de data,
    antes da linha 'Saldo Final'. Recalcula os saldos, os totais de créditos e débitos, a
    posição da conta corrente e o Sumário. Devolve o novo saldo final.

    `debitos_positivos`: escreve o débito sem sinal na coluna de movimentação — o parser
    tem de recuperar o sinal pela diferença de saldo.
    """
    ws = _aba(wb, "Conta Corrente")
    r_ini = _linha(ws, "Saldo Anterior")
    r_cab = r_ini - 1
    cols = _colunas(ws, r_cab)
    c_data, c_desc = _col(cols, "data"), _col(cols, "descricao")
    c_mov, c_saldo = _col(cols, "movimentacao"), _col(cols, "saldo")
    r_cred = _linha(ws, "Total de Creditos", a_partir_de=r_ini)

    saldo_inicial = float(ws.cell(r_ini, c_saldo).value)
    existentes: list[tuple[datetime, str, float]] = []
    saldo_final_linha: tuple[datetime, str, float] | None = None
    for r in range(r_ini + 1, r_cred):
        item = (ws.cell(r, c_data).value, ws.cell(r, c_desc).value, float(ws.cell(r, c_mov).value or 0))
        if _norm(item[1]).startswith("saldo final"):
            saldo_final_linha = item
        else:
            existentes.append(item)
    saldo_final_antigo = float(ws.cell(r_cred - 1, c_saldo).value)

    novos = [(_dt(d), desc, float(v)) for d, desc, v in lancamentos]
    linhas = sorted(existentes + novos, key=lambda i: i[0])   # estável: existentes primeiro
    if saldo_final_linha:
        linhas.append(saldo_final_linha)

    a_inserir = len(linhas) - (r_cred - r_ini - 1)
    if a_inserir > 0:
        ws.insert_rows(r_cred, amount=a_inserir)
    r_cred += max(a_inserir, 0)

    saldo = saldo_inicial
    for i, (d, desc, v) in enumerate(linhas):
        r = r_ini + 1 + i
        saldo = round(saldo + v, 2)
        ws.cell(r, c_data).value = d
        ws.cell(r, c_desc).value = desc
        ws.cell(r, c_mov).value = abs(v) if (debitos_positivos and v < 0) else v
        ws.cell(r, c_saldo).value = saldo

    creditos = round(sum(v for _, _, v in linhas if v > 0), 2)
    debitos = round(sum(v for _, _, v in linhas if v < 0), 2)
    r_deb = _linha(ws, "Total de Debitos", a_partir_de=r_cred)
    ws.cell(r_cred, c_mov).value = creditos if creditos else "-"
    ws.cell(r_deb, c_mov).value = debitos if debitos else "-"

    # Posição da conta corrente e Sumário acompanham o novo saldo
    r_pos = _linha(ws, "Data")
    c_valor = _col(_colunas(ws, r_pos), "valor financeiro")
    ws.cell(r_pos + 1, c_valor).value = saldo
    _somar_no_sumario(wb, "Conta Corrente", round(saldo - saldo_final_antigo, 2))
    return saldo


# ---------------------------------------------------------------------------
# Movimentação de renda variável
# ---------------------------------------------------------------------------

def adicionar_movimentacao_rv(wb: Workbook, bloco: str, linhas: list[dict]) -> None:
    """
    Acrescenta linhas ao bloco 'Movimentação > {bloco}' (ex.: 'Ações', 'Fundos Listados').
    Cada linha: data, transacao, codigo, qtde, preco, valor_bruto, corretagem,
    valor_liquido e, em Fundos Listados, tipo. Atualiza os totais de compras e vendas.
    Não mexe nas posições — o teste decide se quer coerência com elas.
    """
    ws = _aba(wb, "Renda Variavel")
    r_titulo = _linha(ws, f"Movimentacao > {bloco}")
    r_cab = _linha(ws, "Data", a_partir_de=r_titulo)
    cols = _colunas(ws, r_cab)
    r_compras = _linha(ws, "Total de Compras", a_partir_de=r_cab)
    ws.insert_rows(r_compras, amount=len(linhas))
    r_compras += len(linhas)

    campos = {
        "data": "data", "transacao": "transacao", "codigo": "codigo", "tipo": "tipo",
        "qtde": "qtde", "preco": "preco", "valor_bruto": "valor bruto",
        "corretagem": "corretagem", "valor_liquido": "valor liquido",
    }
    for i, linha in enumerate(linhas):
        r = r_compras - len(linhas) + i
        for chave, cabecalho in campos.items():
            if chave not in linha:
                continue
            valor = linha[chave]
            ws.cell(r, _col(cols, cabecalho)).value = _dt(valor) if isinstance(valor, date) else valor

    for rotulo, operacao in (("Total de Compras", "compra"), ("Total de Vendas", "venda")):
        r_tot = _linha(ws, rotulo, a_partir_de=r_cab)
        do_tipo = [l for l in _linhas_rv(ws, r_cab, r_compras, cols) if _norm(l["transacao"]).startswith(operacao)]
        for chave, cabecalho in (("valor_bruto", "valor bruto"), ("corretagem", "corretagem"),
                                 ("valor_liquido", "valor liquido")):
            soma = round(sum(abs(l.get(chave) or 0) for l in do_tipo), 2)
            ws.cell(r_tot, _col(cols, cabecalho)).value = soma if soma else "-"


def _linhas_rv(ws, r_cab: int, r_fim: int, cols: dict[str, int]) -> list[dict]:
    out = []
    for r in range(r_cab + 1, r_fim):
        transacao = ws.cell(r, _col(cols, "transacao")).value
        if not transacao:
            continue
        out.append({
            "transacao": transacao,
            "valor_bruto": ws.cell(r, _col(cols, "valor bruto")).value,
            "corretagem": ws.cell(r, _col(cols, "corretagem")).value,
            "valor_liquido": ws.cell(r, _col(cols, "valor liquido")).value,
        })
    for item in out:
        for k in ("valor_bruto", "corretagem", "valor_liquido"):
            if not isinstance(item[k], (int, float)):
                item[k] = 0.0
    return out


# ---------------------------------------------------------------------------
# Lotes de renda fixa
# ---------------------------------------------------------------------------

def adicionar_lote_rf(wb: Workbook, sigla: str, lote: dict) -> None:
    """
    Acrescenta um lote ao 'Detalhamento > TESOURO DIRETO - {sigla}'. Campos do lote:
    ativo, emissao, vencimento, aquisicao, taxa_compra, quantidade, preco_compra,
    valor_compra, preco, saldo_bruto. Não mexe na posição consolidada.
    """
    ws = _aba(wb, "Renda Fixa")
    r_titulo = _linha(ws, f"Detalhamento > TESOURO DIRETO - {sigla} |")
    r_cab = _linha(ws, "Ativo", a_partir_de=r_titulo)
    cols = _colunas(ws, r_cab)
    r_total = _linha(ws, "Total", a_partir_de=r_cab + 1, coluna=2)
    ws.insert_rows(r_total, amount=1)

    campos = {
        "ativo": "ativo", "emissao": "emissao", "vencimento": "vencimento",
        "aquisicao": "aquisicao", "taxa_compra": "taxa compra", "quantidade": "quantidade",
        "preco_compra": "preco compra", "valor_compra": "valor compra", "preco": "preco r$",
        "saldo_bruto": "saldo bruto",
    }
    for chave, cabecalho in campos.items():
        if chave in lote:
            valor = lote[chave]
            ws.cell(r_total, _col(cols, cabecalho)).value = _dt(valor) if isinstance(valor, date) else valor
